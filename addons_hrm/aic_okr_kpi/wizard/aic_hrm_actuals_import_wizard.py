# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
import base64
import csv
import datetime
import io

from odoo import fields, models, _
from odoo.exceptions import UserError

REQUIRED_COLUMNS = ('code', 'value')
KNOWN_COLUMNS = ('type', 'code', 'employee', 'date_from', 'date_to',
                 'value', 'note')
COLUMN_ALIASES = {'date': 'date_from', 'kpi': 'code', 'kr': 'code'}


class AicHrmActualsImportWizard(models.TransientModel):
    """Bulk-load actual values exported from another system.

    One flat table (CSV or first Excel sheet), one row per measurement:

        type | code | employee | date_from | date_to | value | note

    ``type`` is ``kpi`` (default) or ``kr``. KPI rows upsert period
    results with source 'import'; KR rows land as check-ins so the
    scoring chain reacts normally. Rows never touch confirmed periods
    and never overwrite manual entries - manual always wins.
    """
    _name = 'aic.hrm.actuals.import.wizard'
    _description = 'Actuals Import (Excel/CSV)'

    cycle_id = fields.Many2one('aic.hrm.cycle', required=True)
    file = fields.Binary(required=True)
    filename = fields.Char()
    state = fields.Selection([
        ('upload', 'Upload'),
        ('preview', 'Preview'),
        ('done', 'Done'),
    ], default='upload')
    preview = fields.Text(readonly=True)
    warning_log = fields.Text(readonly=True)
    imported_period_count = fields.Integer(readonly=True)
    imported_checkin_count = fields.Integer(readonly=True)

    # ---- parsing ----

    def _raw_rows(self):
        """Return a list of row tuples including the header row."""
        data = base64.b64decode(self.file)
        name = (self.filename or '').lower()
        if name.endswith('.xlsx'):
            try:
                import openpyxl
            except ImportError as error:
                raise UserError(_(
                    "The Python library 'openpyxl' is required to read "
                    "Excel files. Ask your administrator to install it "
                    "(pip install openpyxl).")) from error
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), data_only=True, read_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            return [tuple(row) for row in sheet.iter_rows(values_only=True)]
        text = data.decode('utf-8-sig', errors='replace')
        sample = text[:2048]
        delimiter = ';' if sample.count(';') > sample.count(',') else ','
        return [tuple(row) for row in
                csv.reader(io.StringIO(text), delimiter=delimiter)]

    @staticmethod
    def _to_date(value):
        if not value:
            return False
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        return fields.Date.to_date(str(value).strip())

    @staticmethod
    def _to_float(value):
        if value in (None, ''):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(' ', '')
        if ',' in text and '.' not in text:
            text = text.replace(',', '.')
        else:
            text = text.replace(',', '')
        return float(text)

    def _parse(self):
        """Return (rows, warnings); rows are normalized dicts."""
        raw = self._raw_rows()
        if not raw:
            raise UserError(_("The file is empty."))
        header = []
        for cell in raw[0]:
            key = str(cell or '').strip().lower()
            header.append(COLUMN_ALIASES.get(key, key))
        missing = [column for column in REQUIRED_COLUMNS
                   if column not in header]
        if missing:
            raise UserError(_(
                "Missing column(s): %(columns)s. Expected a header row "
                "with at least: code, value (and optionally: type, "
                "employee, date_from, date_to, note).",
                columns=', '.join(missing)))
        rows, warnings = [], []
        for line_no, raw_row in enumerate(raw[1:], start=2):
            if not any(cell not in (None, '') for cell in raw_row):
                continue
            entry = dict(zip(header, raw_row))
            code = str(entry.get('code') or '').strip()
            if not code:
                warnings.append(_("Line %s: no code - skipped.") % line_no)
                continue
            try:
                value = self._to_float(entry.get('value'))
            except ValueError:
                value = None
            if value is None:
                warnings.append(_(
                    "Line %(line)s (%(code)s): value is not a number - "
                    "skipped.", line=line_no, code=code))
                continue
            try:
                date_from = self._to_date(entry.get('date_from'))
                date_to = self._to_date(entry.get('date_to'))
            except ValueError:
                warnings.append(_(
                    "Line %(line)s (%(code)s): unreadable date - skipped.",
                    line=line_no, code=code))
                continue
            row_type = str(entry.get('type') or 'kpi').strip().lower()
            if row_type not in ('kpi', 'kr'):
                warnings.append(_(
                    "Line %(line)s (%(code)s): unknown type %(type)s - "
                    "skipped.", line=line_no, code=code, type=row_type))
                continue
            rows.append({
                'line': line_no,
                'type': row_type,
                'code': code,
                'employee': str(entry.get('employee') or '').strip(),
                'date_from': date_from,
                'date_to': date_to or date_from,
                'value': value,
                'note': str(entry.get('note') or '').strip(),
            })
        return rows, warnings

    # ---- matching ----

    def _find_kpi_target(self, row, warnings):
        kpi = self.env['aic.hrm.kpi'].search(
            [('code', '=ilike', row['code'])], limit=1)
        if not kpi:
            warnings.append(_(
                "Line %(line)s: no KPI with code %(code)s.",
                line=row['line'], code=row['code']))
            return self.env['aic.hrm.kpi.target']
        domain = [('kpi_id', '=', kpi.id),
                  ('cycle_id', '=', self.cycle_id.id)]
        if row['employee']:
            employee = self.env['hr.employee'].search(
                [('name', '=ilike', row['employee'])], limit=1)
            if not employee:
                warnings.append(_(
                    "Line %(line)s (%(code)s): no employee named "
                    "%(name)s.", line=row['line'], code=row['code'],
                    name=row['employee']))
                return self.env['aic.hrm.kpi.target']
            domain.append(('employee_id', '=', employee.id))
        else:
            domain.append(('employee_id', '=', False))
        target = self.env['aic.hrm.kpi.target'].search(domain, limit=1)
        if not target:
            warnings.append(_(
                "Line %(line)s (%(code)s): no KPI target in cycle "
                "%(cycle)s for this owner.", line=row['line'],
                code=row['code'], cycle=self.cycle_id.name))
        return target

    def _find_key_result(self, row, warnings):
        kr = self.env['aic.hrm.key.result'].search(
            [('code', '=ilike', row['code']),
             ('objective_id.cycle_id', '=', self.cycle_id.id)], limit=1)
        if not kr:
            warnings.append(_(
                "Line %(line)s: no key result with code %(code)s in "
                "cycle %(cycle)s.", line=row['line'], code=row['code'],
                cycle=self.cycle_id.name))
        return kr

    # ---- actions ----

    def action_preview(self):
        self.ensure_one()
        rows, warnings = self._parse()
        kpi_ok = kr_ok = 0
        for row in rows:
            if row['type'] == 'kpi':
                if not row['date_from']:
                    warnings.append(_(
                        "Line %(line)s (%(code)s): KPI rows need "
                        "date_from.", line=row['line'], code=row['code']))
                    continue
                if self._find_kpi_target(row, warnings):
                    kpi_ok += 1
            else:
                if self._find_key_result(row, warnings):
                    kr_ok += 1
        self.write({
            'state': 'preview',
            'preview': _(
                "%(kpi)s KPI period value(s) and %(kr)s key-result "
                "check-in(s) ready to import into cycle %(cycle)s.",
                kpi=kpi_ok, kr=kr_ok, cycle=self.cycle_id.name),
            'warning_log': '\n'.join(warnings) or False,
        })
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        self.cycle_id.ensure_editable()
        rows, warnings = self._parse()
        Period = self.env['aic.hrm.kpi.period.result']
        periods = checkins = 0
        for row in rows:
            if row['type'] == 'kpi':
                if not row['date_from']:
                    warnings.append(_(
                        "Line %(line)s (%(code)s): KPI rows need "
                        "date_from.", line=row['line'], code=row['code']))
                    continue
                target = self._find_kpi_target(row, warnings)
                if not target:
                    continue
                existing = Period.search(
                    [('kpi_target_id', '=', target.id),
                     ('date_from', '=', row['date_from'])], limit=1)
                if existing:
                    if existing.state == 'confirmed':
                        warnings.append(_(
                            "Line %(line)s (%(code)s): period "
                            "%(date)s is already confirmed - skipped.",
                            line=row['line'], code=row['code'],
                            date=row['date_from']))
                        continue
                    if existing.source == 'manual':
                        warnings.append(_(
                            "Line %(line)s (%(code)s): period %(date)s "
                            "was entered manually - manual wins, "
                            "skipped.", line=row['line'],
                            code=row['code'], date=row['date_from']))
                        continue
                    existing.write({
                        'actual': row['value'],
                        'source': 'import',
                        'note': row['note'] or existing.note,
                    })
                else:
                    Period.create({
                        'kpi_target_id': target.id,
                        'date_from': row['date_from'],
                        'date_to': row['date_to'],
                        'actual': row['value'],
                        'source': 'import',
                        'note': row['note'] or False,
                    })
                periods += 1
            else:
                kr = self._find_key_result(row, warnings)
                if not kr:
                    continue
                self.env['aic.hrm.checkin'].create({
                    'kr_id': kr.id,
                    'date': row['date_from'] or fields.Date.context_today(
                        self),
                    'value_current': row['value'],
                    'note': row['note'] or _(
                        'Imported from %s') % (self.filename or 'file'),
                })
                checkins += 1
        self.write({
            'state': 'done',
            'imported_period_count': periods,
            'imported_checkin_count': checkins,
            'warning_log': '\n'.join(warnings) or False,
        })
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
