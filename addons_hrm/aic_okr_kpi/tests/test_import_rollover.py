# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
import base64
import io

from odoo.tests import tagged

from .common import OkrCase


def build_dlsp_workbook():
    """Fixture mirroring the real 3-sheet customer workbook structure
    (OKR_2026 / KPI_CHI_TIET / PHAN_CONG) with fictional content."""
    import openpyxl
    workbook = openpyxl.Workbook()

    okr = workbook.active
    okr.title = 'OKR_2026'
    okr.append(['Mã Objective', 'Objective 2026', 'Tỷ trọng O (%)', 'Mã KR',
                'Key Result', 'Thước đo', 'Target 2026', 'Đơn vị', 'Chủ trì',
                'Quý trọng tâm', 'Ưu tiên', 'Ghi chú'])
    okr.append(['O1', 'Launch the streaming platform', 60, 'KR1.1',
                'Go-live completed', 'Completion level', 100, '%',
                'Nam Member', 'Q2-Q3', 'Cao', ''])
    okr.append(['', '', 60, 'KR1.2', 'Paying subscribers onboarded',
                'Subscribers', 5000, 'users', 'Mai Manager', 'Q3-Q4',
                'Trung bình', ''])
    okr.append(['O2', 'Standardize data platform', 40, 'KR2.1',
                'Data sources connected', 'Sources', 6, 'nguồn',
                'Unknown Person', 'Q1-Q4', 'Cao', ''])

    kpi = workbook.create_sheet('KPI_CHI_TIET')
    kpi.append(['KPI ID', 'Nhóm KPI', 'Objective', 'KPI', 'Chiều đo',
                'Cách tổng hợp', 'Đơn vị', 'Trọng số', 'Target 2026',
                'Chủ trì', 'Nguồn đo', 'Ghi chú'])
    kpi.append(['KPI01', 'Product', 'O1', 'Platform go-live on schedule',
                'Càng cao càng tốt', 'Cuối kỳ', '%', 60, 100,
                'Nam Member', 'Release report', ''])
    kpi.append(['KPI02', 'Product', 'O1', 'Critical bugs per quarter',
                'Càng thấp càng tốt', 'Bình quân', 'lỗi/quý', 20, 3,
                'Nam Member', 'Bug tracker', ''])
    kpi.append(['KPI03', 'Data', 'O2', 'Integration standard published',
                'Đạt/Không đạt', 'Cuối kỳ', '0/1', 20, 1,
                'Mai Manager', 'Decision record', ''])

    assign = workbook.create_sheet('PHAN_CONG')
    assign.append(['TT', 'Họ và tên', 'Vị trí/Nhóm nhân sự',
                   'Trách nhiệm trọng tâm', 'Objective liên quan',
                   'KPI đề xuất giao', 'Tổng trọng số cá nhân đề xuất (%)'])
    assign.append([1, 'Nam Member', 'Product Owner', 'Ship the platform',
                   'O1', 'KPI01,KPI02', 100])
    assign.append([2, 'Mai Manager', 'Head of Department', 'Run the org',
                   'O1,O2', 'KPI03', 100])
    assign.append([3, 'Missing Person', 'Data Engineer', 'Pipelines',
                   'O2', 'KPI03', 100])
    assign.append([4, 'Duo One, Duo Two', 'Shared squad', 'Joint delivery',
                   'O1,O2', 'KPI02,KPI03', 100])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return base64.b64encode(buffer.getvalue())


@tagged('post_install', '-at_install', 'aic_okr_kpi')
class TestExcelImport(OkrCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Wizard = cls.env['aic.hrm.import.wizard']

    def _make_wizard(self, **kw):
        vals = {
            'cycle_id': self.year.id,
            'file': build_dlsp_workbook(),
            'filename': 'fixture.xlsx',
        }
        vals.update(kw)
        return self.Wizard.create(vals)

    def test_preview_reports_structure(self):
        wizard = self._make_wizard()
        wizard.action_preview()
        self.assertEqual(wizard.state, 'preview')
        self.assertIn('O1', wizard.preview)
        self.assertIn('KPI01', wizard.preview)
        # unmatched employee names surface as warnings, not silent failures
        self.assertIn('Unknown Person', wizard.warning_log)
        self.assertIn('Missing Person', wizard.warning_log)

    def test_import_creates_okr_kpi_assignment(self):
        wizard = self._make_wizard()
        wizard.action_preview()
        wizard.action_import()
        objectives = self.Objective.search(
            [('cycle_id', '=', self.year.id), ('code', 'in', ['O1', 'O2'])])
        self.assertEqual(len(objectives), 2)
        o1 = objectives.filtered(lambda o: o.code == 'O1')
        self.assertAlmostEqual(o1.weight, 60.0)
        self.assertEqual(len(o1.kr_ids), 2)
        kr11 = o1.kr_ids.filtered(lambda kr: kr.code == 'KR1.1')
        self.assertEqual(kr11.employee_id, self.member_employee)

        kpis = self.env['aic.hrm.kpi'].search(
            [('code', 'in', ['KPI01', 'KPI02', 'KPI03'])])
        self.assertEqual(len(kpis), 3)
        bug_kpi = kpis.filtered(lambda k: k.code == 'KPI02')
        self.assertEqual(bug_kpi.direction, 'lower')
        self.assertEqual(bug_kpi.aggregation, 'average')
        pass_kpi = kpis.filtered(lambda k: k.code == 'KPI03')
        self.assertEqual(pass_kpi.direction, 'boolean')

        assignments = self.env['aic.hrm.kpi.assignment'].search(
            [('cycle_id', '=', self.year.id)])
        self.assertEqual(len(assignments), 2,
                         "rows with unmatched employees are skipped")
        member_card = assignments.filtered(
            lambda a: a.employee_id == self.member_employee)
        self.assertEqual(len(member_card.line_ids), 2)
        # The sheet weights are each KPI's share of the DEPARTMENT plan,
        # so this person's two lines carried 60 + 20. A scorecard has to
        # total 100 before it can be submitted, so the import scales them
        # while keeping the 3:1 proportion the planner expressed.
        self.assertAlmostEqual(member_card.total_weight, 100.0,
                               msg="a scorecard must be submittable")

    def test_import_is_idempotent_on_rerun(self):
        wizard = self._make_wizard()
        wizard.action_preview()
        wizard.action_import()
        wizard2 = self._make_wizard()
        wizard2.action_preview()
        wizard2.action_import()
        self.assertEqual(self.Objective.search_count(
            [('cycle_id', '=', self.year.id), ('code', '=', 'O1')]), 1)

    def test_scorecard_weights_are_scaled_to_100(self):
        """Imported weights are the KPI's share of the DEPARTMENT plan, so
        one person's lines summed to whatever slice they happened to hold.
        A scorecard cannot be submitted below 100, so every imported
        scorecard was stuck in draft - the import produced data the
        product refuses."""
        wizard = self._make_wizard()
        wizard.action_preview()
        wizard.action_import()
        assignments = self.env['aic.hrm.kpi.assignment'].search(
            [('cycle_id', '=', self.year.id)])
        self.assertTrue(assignments)
        for assignment in assignments:
            total = sum(assignment.line_ids.mapped('weight'))
            self.assertAlmostEqual(
                total, 100.0, places=2,
                msg='%s totals %s' % (assignment.employee_id.name, total))

    def test_scaling_keeps_relative_proportions(self):
        """Scaling must not flatten the planner's intent: a KPI weighted
        three times another stays three times another."""
        wizard = self._make_wizard()
        wizard.action_preview()
        wizard.action_import()
        member = self.env['hr.employee'].search(
            [('name', '=', 'Nam Member')], limit=1)
        assignment = self.env['aic.hrm.kpi.assignment'].search(
            [('cycle_id', '=', self.year.id),
             ('employee_id', '=', member.id)], limit=1)
        weights = sorted(assignment.line_ids.mapped('weight'), reverse=True)
        self.assertEqual(len(weights), 2)
        # KPI01 weighed 60 against KPI02's 20 in the sheet.
        self.assertAlmostEqual(weights[0] / weights[1], 3.0, places=1)

    def test_row_naming_several_people_creates_one_each(self):
        """A cell reading "Duo One, Duo Two" used to create an employee of
        that name - a person who does not exist, holding a scorecard
        nobody owns."""
        wizard = self._make_wizard(create_missing_employees=True)
        wizard.action_preview()
        wizard.action_import()
        Employee = self.env['hr.employee']
        self.assertFalse(
            Employee.search([('name', '=', 'Duo One, Duo Two')]),
            'the merged name was created as a person')
        for name in ('Duo One', 'Duo Two'):
            person = Employee.search([('name', '=', name)], limit=1)
            self.assertTrue(person, '%s was not created' % name)
            assignment = self.env['aic.hrm.kpi.assignment'].search(
                [('cycle_id', '=', self.year.id),
                 ('employee_id', '=', person.id)], limit=1)
            self.assertTrue(assignment, '%s has no scorecard' % name)
            self.assertAlmostEqual(
                sum(assignment.line_ids.mapped('weight')), 100.0, places=2)

    def test_broken_file_reports_error(self):
        wizard = self._make_wizard(file=base64.b64encode(b'not an xlsx'))
        wizard.action_preview()
        self.assertEqual(wizard.state, 'error')
        self.assertTrue(wizard.error_log)


@tagged('post_install', '-at_install', 'aic_okr_kpi')
class TestRollover(OkrCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.objective = cls._make_objective(
            employee_id=cls.member_employee.id)
        cls.kr = cls._make_kr(cls.objective, current=80.0)
        cls.kpi = cls.env['aic.hrm.kpi'].create({
            'name': 'Roll KPI', 'code': 'KPI-ROLL'})
        cls.target = cls.env['aic.hrm.kpi.target'].create({
            'kpi_id': cls.kpi.id, 'cycle_id': cls.year.id,
            'employee_id': cls.member_employee.id, 'target_value': 100.0,
        })
        cls.env['aic.hrm.kpi.period.result'].create({
            'kpi_target_id': cls.target.id, 'date_from': '2026-01-01',
            'date_to': '2026-01-31', 'actual': 70.0, 'state': 'confirmed'})

    def test_rollover_copies_structure_resets_actuals(self):
        wizard = self.env['aic.hrm.rollover.wizard'].create({
            'source_cycle_id': self.year.id,
            'target_cycle_id': self.other_cycle.id,
        })
        wizard.action_rollover()
        new_objectives = self.Objective.search(
            [('cycle_id', '=', self.other_cycle.id)])
        self.assertEqual(len(new_objectives), 1)
        new_kr = new_objectives.kr_ids
        self.assertEqual(len(new_kr), 1)
        self.assertAlmostEqual(new_kr.current, 0.0, msg="actuals reset")
        self.assertAlmostEqual(new_kr.target, self.kr.target,
                               msg="targets kept, editable afterwards")
        self.assertEqual(new_objectives.state, 'draft')
        new_targets = self.env['aic.hrm.kpi.target'].search(
            [('cycle_id', '=', self.other_cycle.id)])
        self.assertEqual(len(new_targets), 1)
        self.assertFalse(new_targets.period_result_ids,
                         "period results never roll over")

    def test_rollover_remaps_alignment(self):
        child = self._make_objective(
            name='Child objective', parent_id=self.objective.id, weight=5.0)
        wizard = self.env['aic.hrm.rollover.wizard'].create({
            'source_cycle_id': self.year.id,
            'target_cycle_id': self.other_cycle.id,
        })
        wizard.action_rollover()
        new_child = self.Objective.search([
            ('cycle_id', '=', self.other_cycle.id),
            ('name', '=', 'Child objective')])
        self.assertTrue(new_child.parent_id)
        self.assertEqual(new_child.parent_id.cycle_id, self.other_cycle,
                         "parent link remapped inside the new cycle")


def build_grouped_workbook(kr_code='O1.KR1'):
    """A monthly assignment sheet: KPIs in weighted groups, each serving a
    quarterly key result, one milestone target described in words."""
    import openpyxl
    workbook = openpyxl.Workbook()
    okr = workbook.active
    okr.title = 'OKR_2026'
    okr.append(['Mã Objective', 'Objective', 'Tỷ trọng O (%)', 'Mã KR', 'Key Result', 'Thước đo',
                'Target', 'Đơn vị', 'Chủ trì', 'Quý trọng tâm', 'Ưu tiên', 'Ghi chú'])

    kpi = workbook.create_sheet('KPI_CHI_TIET')
    kpi.append(['KPI ID', 'Nhóm KPI', 'Objective', 'KPI', 'Chiều đo', 'Cách tổng hợp', 'Đơn vị',
                'Trọng số', 'Target', 'Chủ trì', 'Nguồn đo', 'Ghi chú',
                'Trọng số nhóm (%)', 'Trọng số trong nhóm (%)', 'Mã KR', 'Chỉ tiêu (nguyên văn)'])
    kpi.append(['KDDV.CV.B1.1', 'B.I KPI Doanh thu', '', 'Channel revenue', 'Càng cao càng tốt',
                'Cuối kỳ', 'tỷ VNĐ', 64, 21.92, 'Nam Member', 'Contract system', '',
                80, 80, kr_code, '≥ 21,92 tỷ (gốc 20,92 + 1,00 bổ sung)'])
    kpi.append(['KDDV.CV.B1.2', 'B.I KPI Doanh thu', '', 'Merchant onboarding', 'Đạt/Không đạt',
                'Cuối kỳ', 'merchant', 16, 1, 'Nam Member', 'E-commerce report', '',
                80, 20, kr_code, 'Ký cam kết ≥ 8 merchant lũy kế'])
    kpi.append(['KDDV.CV.B2.1', 'B.II KPI Quản trị', '', 'Copyright incidents', 'Càng thấp càng tốt',
                'Cuối kỳ', 'vụ', 20, 0, 'Nam Member', 'Legal log', '',
                20, 100, '', '0'])

    assign = workbook.create_sheet('PHAN_CONG')
    assign.append(['TT', 'Họ và tên', 'Vị trí', 'Trách nhiệm', 'Objective', 'KPI', 'Tổng (%)'])
    assign.append([1, 'Nam Member, Mai Manager', 'Chuyên viên KD', 'Kênh', '',
                   'KDDV.CV.B1.1,KDDV.CV.B1.2,KDDV.CV.B2.1', 100])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return base64.b64encode(buffer.getvalue())


@tagged('post_install', '-at_install', 'aic_okr_kpi')
class TestGroupedImport(OkrCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.month = cls.Cycle.create({
            'name': 'May 2026', 'code': 'OKR-M05-GRP', 'cycle_type': 'month',
            'date_start': '2026-05-01', 'date_end': '2026-05-31',
            'parent_id': cls.quarter.id,
        })
        cls.objective = cls._make_objective(cycle_id=cls.quarter.id, code='O1')
        cls.kr = cls._make_kr(cls.objective, code='O1.KR1')

    def _import(self, **kw):
        wizard = self.env['aic.hrm.import.wizard'].create({
            'cycle_id': self.month.id, 'file': build_grouped_workbook(**kw), 'filename': 'grouped.xlsx'})
        wizard.action_preview()
        wizard.action_import()
        return wizard

    def test_groups_weights_and_key_results_are_imported(self):
        self._import()
        cards = self.env['aic.hrm.kpi.assignment'].search([('cycle_id', '=', self.month.id)])
        self.assertEqual(len(cards), 2, 'one scorecard per person named on the row')
        for card in cards:
            self.assertEqual(sorted(card.group_ids.mapped('weight')), [20.0, 80.0])
            by_code = {line.kpi_target_id.kpi_id.code: line for line in card.line_ids}
            self.assertEqual(by_code['KDDV.CV.B1.1'].weight_in_group, 80.0)
            self.assertAlmostEqual(by_code['KDDV.CV.B1.1'].weight, 64.0)
            self.assertAlmostEqual(by_code['KDDV.CV.B1.2'].weight, 16.0)
            self.assertAlmostEqual(by_code['KDDV.CV.B2.1'].weight, 20.0)
            self.assertTrue(card.weight_ok)
            revenue = by_code['KDDV.CV.B1.1'].kpi_target_id
            self.assertEqual(revenue.kr_id, self.kr)
            self.assertEqual(revenue.objective_id, self.objective)
            self.assertEqual(revenue.target_note, '≥ 21,92 tỷ (gốc 20,92 + 1,00 bổ sung)')
            incidents = by_code['KDDV.CV.B2.1'].kpi_target_id
            self.assertEqual((incidents.direction, incidents.target_value), ('lower', 0.0))
            self.assertFalse(incidents.kr_id)

    def test_reimport_changes_nothing(self):
        self._import()
        before = self.env['aic.hrm.kpi.assignment.line'].search_count([('assignment_id.cycle_id', '=', self.month.id)])
        self._import()
        after = self.env['aic.hrm.kpi.assignment.line'].search_count([('assignment_id.cycle_id', '=', self.month.id)])
        self.assertEqual(before, after)
        groups = self.env['aic.hrm.kpi.assignment.group'].search_count([('assignment_id.cycle_id', '=', self.month.id)])
        self.assertEqual(groups, 4)

    def test_unknown_key_result_is_reported_not_guessed(self):
        wizard = self._import(kr_code='O9.KR9')
        self.assertIn('O9.KR9', wizard.warning_log)
        target = self.env['aic.hrm.kpi.target'].search([
            ('cycle_id', '=', self.month.id), ('kpi_id.code', '=', 'KDDV.CV.B1.1')], limit=1)
        self.assertFalse(target.kr_id)
