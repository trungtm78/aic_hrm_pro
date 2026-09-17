# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""Read the Sales & Services actual revenue and cost workbooks into one dataset.

The customer keeps two spreadsheets outside Odoo:

* the receipts register (one sheet, a block of seven columns per month:
  invoiced revenue before VAT / VAT / after VAT, revenue delivered but not yet
  invoiced before VAT / VAT / after VAT, note), grouped by business line;
* the cost ledger (one sheet per month, debit postings by account, the Roman
  numbered rows carrying each account group's total).

Decisions agreed with the customer's owner on 2026-09-17:

* actual revenue = invoiced + delivered-not-invoiced, before VAT, in the month
  column where the register puts it;
* the register's business lines map onto the five planned revenue streams as
  SECTION_STREAM says; the digital-services stream has no line and is 0;
* costs are tracked for the department, never weighted into anyone's score;
* only numbers present in the files are produced; nothing is estimated.

The script also lists what looks irregular (VAT not 8%, negative amounts, a
partner both invoiced and not invoiced in one month, lump invoices) so the
irregularities travel with the numbers instead of being silently scored.

Nothing here talks to Odoo.

Run:  python tools/extract_kddv_actuals.py [--source Docs/OKR] [--out uat/data/kddv_actuals_q3_2026.json]
"""
import argparse
import json
import pathlib
import statistics

_REPO = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = _REPO / 'Docs' / 'OKR'
DEFAULT_OUT = _REPO / 'uat' / 'data' / 'kddv_actuals_q3_2026.json'

REVENUE_FILE = 'Template_PhieuThu_2026_20260917.xlsx'
COST_FILE = 'chi_phi_2026.xlsx'
YEAR = 2026
QUARTER_MONTHS = (7, 8, 9)
BILLION = 1e9
VAT_RATE = 0.08
VAT_TOLERANCE = 0.005
LUMP_FACTOR = 2.5
# Below this (VND) a cell is too small for its VAT ratio or timing to matter.
MATERIAL = 1e7

# Register column layout: data starts at column C, seven columns per month.
FIRST_MONTH_COLUMN = 2
MONTH_WIDTH = 7
TOTAL_COLUMN = FIRST_MONTH_COLUMN + 12 * MONTH_WIDTH + 6  # "Tổng doanh thu" before VAT
FIRST_DATA_ROW = 4

STREAMS = {
    'telco': 'Tiếp phát sóng kênh (Telco/ISP)',
    'vtvshop_mg': 'VTVshop MG',
    'tnnd': 'DV trải nghiệm nội dung (TNND)',
    'fast': 'FAST Channel & Chuyên trang',
    'digital': 'Dịch vụ số',
}
SECTION_STREAM = {
    'I': 'tnnd',        # Dịch vụ khác (app-store and cable resellers of the packages)
    'II': 'telco',      # Doanh thu cấp quyền tiếp phát sóng các kênh
    'III': 'fast',      # Doanh thu từ Chuyên trang trên VTVgo
    'IV': 'vtvshop_mg', # Doanh thu từ VTVshop và Dịch vụ GTGT
    'V': 'tnnd',        # Gói trải nghiệm VTVGo
    'VI': 'tnnd',       # Gói VTV Thể thao
    'VII': 'tnnd',      # Gói World Cup
}

# Scorecard lines whose monthly target is a revenue figure, and the streams
# that make up their actual. Codes are the ones on production (KDDV.<pos>.<line>).
KPI_STREAMS = {
    ('TP-KD', 'B1.1'): ('telco', 'vtvshop_mg', 'tnnd', 'fast', 'digital'),
    ('TP-KD', 'B1.2'): ('telco',),
    ('TP-KD', 'B1.3'): ('vtvshop_mg',),
    ('TP-KD', 'B1.4'): ('tnnd', 'fast', 'digital'),
    ('PPT-KD', 'B1.1'): ('tnnd',),
    ('PPT-KD', 'B1.2'): ('fast',),
    ('CV-KD1', 'B1.1'): ('telco',),
    ('CV-KD1', 'B1.2'): ('vtvshop_mg',),
    ('CV-KD2', 'B1.1'): ('tnnd',),
    ('CV-KD2', 'B1.3'): ('fast',),
    ('CV-KD2', 'B2.3'): ('fast',),
}
REVENUE_KR = 'O1.KR1'

TRACKING = {
    'cost': {'code': 'KDDV.PHONG.CP', 'name': 'Chi phí hoạt động (sổ kế toán)', 'unit': 'tỷ VNĐ',
             'direction': 'lower'},
    'gross_profit': {'code': 'KDDV.PHONG.LNG', 'name': 'Lợi nhuận gộp (Doanh thu − Chi phí)',
                     'unit': 'tỷ VNĐ', 'direction': 'higher'},
    'margin': {'code': 'KDDV.PHONG.TSLN', 'name': 'Tỷ suất lợi nhuận gộp', 'unit': '%',
               'direction': 'higher'},
}
COST_SOURCE = 'Sổ kế toán, phát sinh nợ TK 242/622/627/635/642 — chi_phi_2026.xlsx'


def amount(value):
    return float(value) if isinstance(value, (int, float)) else 0.0


def billions(value):
    return round(value / BILLION, 6)


def fmt(value):
    """Vietnamese number style, 2 decimals: 1.234,56"""
    text = '{:,.2f}'.format(value)
    return text.replace(',', ' ').replace('.', ',').replace(' ', '.')


def is_section(cell):
    return isinstance(cell, str) and cell.strip() in SECTION_STREAM


def read_register(path):
    """Partner rows: section, partner, 12 months of invoiced/uninvoiced (+VAT), note."""
    import openpyxl
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    partners, section = [], None
    for number, row in enumerate(sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), FIRST_DATA_ROW):
        if is_section(row[0]):
            section = row[0].strip()
            continue
        if not row[1] or section is None:
            continue
        months = {}
        for month in range(1, 13):
            base = FIRST_MONTH_COLUMN + (month - 1) * MONTH_WIDTH
            months[month] = {
                'invoiced': amount(row[base]), 'invoiced_vat': amount(row[base + 1]),
                'uninvoiced': amount(row[base + 3]), 'uninvoiced_vat': amount(row[base + 4]),
                'note': (row[base + 6] or '').strip() if isinstance(row[base + 6], str) else '',
            }
        partners.append({'row': number, 'section': section, 'stream': SECTION_STREAM[section],
                         'partner': ' '.join(str(row[1]).split()), 'months': months,
                         'year_total': amount(row[TOTAL_COLUMN])})
    return partners


def read_ledger(path):
    """{month: {'total': x, 'groups': [{account, name, amount}]}} for months with postings."""
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=True)
    ledger = {}
    for month in range(1, 13):
        sheet = book['Tháng %d' % month]
        groups = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            label = row[0]
            if isinstance(label, str) and label.strip().isalpha() and row[2]:
                groups.append({'account': str(row[1]), 'name': ' '.join(str(row[3]).split()),
                               'amount': amount(row[2])})
        if groups:
            ledger[month] = {'total': sum(group['amount'] for group in groups), 'groups': groups}
    return ledger


def irregularities(partners, months):
    found = []
    for partner in partners:
        label = '%s (mục %s, %s)' % (partner['partner'], partner['section'], STREAMS[partner['stream']])
        values = [m['invoiced'] + m['uninvoiced'] for m in partner['months'].values()
                  if m['invoiced'] + m['uninvoiced'] > 0]
        typical = statistics.median(values) if len(values) >= 3 else None
        for month in months:
            cell = partner['months'][month]
            for kind, vat_key, word in (('invoiced', 'invoiced_vat', 'đã xuất HĐ'),
                                        ('uninvoiced', 'uninvoiced_vat', 'chưa xuất HĐ')):
                base = cell[kind]
                if base < 0:
                    found.append({'month': month, 'stream': partner['stream'], 'partner': partner['partner'],
                                  'kind': 'negative',
                                  'text': 'T%d %s: doanh thu %s âm %s tỷ (điều chỉnh giảm)'
                                          % (month, label, word, fmt(billions(base)))})
                if abs(base) >= MATERIAL and abs(cell[vat_key] / base - VAT_RATE) > VAT_TOLERANCE:
                    found.append({'month': month, 'stream': partner['stream'], 'partner': partner['partner'],
                                  'kind': 'vat',
                                  'text': 'T%d %s: VAT %s = %.1f%% (không phải 8%%)'
                                          % (month, label, word, 100 * cell[vat_key] / base)})
            if cell['invoiced'] >= MATERIAL and cell['uninvoiced'] >= MATERIAL:
                found.append({'month': month, 'stream': partner['stream'], 'partner': partner['partner'],
                              'kind': 'both',
                              'text': 'T%d %s: vừa đã xuất HĐ %s tỷ vừa chưa xuất HĐ %s tỷ trong cùng tháng '
                                      '(cần xác nhận phần chưa HĐ có thuộc tháng khác không)'
                                      % (month, label, fmt(billions(cell['invoiced'])),
                                         fmt(billions(cell['uninvoiced'])))})
            total = cell['invoiced'] + cell['uninvoiced']
            if typical and total >= LUMP_FACTOR * typical and total >= BILLION:
                found.append({'month': month, 'stream': partner['stream'], 'partner': partner['partner'],
                              'kind': 'lump',
                              'text': 'T%d %s: %s tỷ, gấp %.1f lần mức thường của đối tác (nghi hoá đơn gộp nhiều tháng)'
                                      % (month, label, fmt(billions(total)), total / typical)})
    return found


def stream_totals(partners, month):
    streams = {key: {'invoiced': 0.0, 'uninvoiced': 0.0, 'total': 0.0, 'partners': []} for key in STREAMS}
    for partner in partners:
        cell = partner['months'][month]
        if not (cell['invoiced'] or cell['uninvoiced']):
            continue
        stream = streams[partner['stream']]
        stream['invoiced'] += billions(cell['invoiced'])
        stream['uninvoiced'] += billions(cell['uninvoiced'])
        stream['partners'].append({'partner': partner['partner'], 'section': partner['section'],
                                   'invoiced': billions(cell['invoiced']),
                                   'uninvoiced': billions(cell['uninvoiced'])})
    for stream in streams.values():
        stream['invoiced'] = round(stream['invoiced'], 6)
        stream['uninvoiced'] = round(stream['uninvoiced'], 6)
        stream['total'] = round(stream['invoiced'] + stream['uninvoiced'], 6)
    return streams


def month_has_revenue(partners, month):
    return any(p['months'][month]['invoiced'] or p['months'][month]['uninvoiced'] for p in partners)


def breakdown_note(streams, keys, month, issues):
    parts = []
    for key in keys:
        stream = streams[key]
        parts.append('%s %s (HĐ %s + chưa HĐ %s)' % (STREAMS[key], fmt(stream['total']),
                                                   fmt(stream['invoiced']), fmt(stream['uninvoiced'])))
    note = 'Thực tế T%d/%d, chưa VAT, tỷ VNĐ: %s. Nguồn: %s.' % (month, YEAR, '; '.join(parts), REVENUE_FILE)
    related = [issue['text'] for issue in issues if issue['month'] == month and issue['stream'] in keys]
    if related:
        note += ' Lưu ý: ' + ' | '.join(related)
    return note


def build(source):
    source = pathlib.Path(source)
    partners = read_register(source / REVENUE_FILE)
    ledger = read_ledger(source / COST_FILE)
    months = [m for m in QUARTER_MONTHS if month_has_revenue(partners, m)]
    issues = irregularities(partners, QUARTER_MONTHS)

    revenue = {}
    for month in QUARTER_MONTHS:
        streams = stream_totals(partners, month)
        revenue[month] = {'has_data': month in months, 'streams': streams,
                          'total': round(sum(s['total'] for s in streams.values()), 6)}

    kpi_actuals = []
    for month in months:
        streams = revenue[month]['streams']
        for (position, line), keys in KPI_STREAMS.items():
            kpi_actuals.append({
                'month': month, 'position': position, 'line': line,
                'code': 'KDDV.%s.%s' % (position, line),
                'value': round(sum(streams[key]['total'] for key in keys), 6),
                'note': breakdown_note(streams, keys, month, issues),
            })

    quarter_total = round(sum(revenue[m]['total'] for m in months), 6)
    kr_checkins = [{
        'code': REVENUE_KR, 'value': quarter_total,
        'note': 'Doanh thu Quý III/%d lũy kế đến hết T%d (chưa VAT, HĐ + chưa HĐ): %s tỷ = %s. '
                'Số liệu có hoá đơn gộp và phần chưa HĐ cần xác nhận — xem ghi chú KPI tháng.'
                % (YEAR, max(months), fmt(quarter_total),
                   ' + '.join('T%d %s' % (m, fmt(revenue[m]['total'])) for m in months)),
    }] if months else []

    tracking = []
    for month in QUARTER_MONTHS:
        cost = ledger.get(month)
        values = {'cost': None, 'gross_profit': None, 'margin': None}
        if cost:
            values['cost'] = billions(cost['total'])
            if revenue[month]['has_data']:
                values['gross_profit'] = round(revenue[month]['total'] - values['cost'], 6)
                values['margin'] = round(100 * values['gross_profit'] / revenue[month]['total'], 4)
        for key, value in values.items():
            if value is None:
                note = 'T%d/%d: chưa có số liệu %s.' % (month, YEAR,
                                                        'chi phí' if not cost else 'doanh thu')
            elif key == 'cost':
                note = 'T%d/%d: %s. %s.' % (month, YEAR, '; '.join(
                    'TK %s %s %s' % (g['account'], g['name'], fmt(billions(g['amount']))) for g in cost['groups']),
                    COST_SOURCE)
            else:
                note = 'T%d/%d: Doanh thu %s − Chi phí %s (tỷ VNĐ). Chi phí là số sổ kế toán, cần xác nhận phạm vi phòng.' % (
                    month, YEAR, fmt(revenue[month]['total']), fmt(values['cost']))
            tracking.append(dict(TRACKING[key], key=key, month=month, value=value, note=note))

    return {
        'year': YEAR, 'months_with_revenue': months,
        'months_with_cost': sorted(ledger),
        'streams': STREAMS,
        'revenue': {str(m): v for m, v in revenue.items()},
        'cost': {str(m): {'total': billions(v['total']),
                          'groups': [dict(g, amount=billions(g['amount'])) for g in v['groups']]}
                 for m, v in ledger.items()},
        'irregularities': issues,
        'register_check': [{'partner': p['partner'], 'row': p['row']} for p in partners
                           if abs(sum(c['invoiced'] + c['uninvoiced'] for c in p['months'].values())
                                  - p['year_total']) > 5],
        'kpi_actuals': kpi_actuals,
        'kr_checkins': kr_checkins,
        'tracking': tracking,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--source', default=str(DEFAULT_SOURCE))
    parser.add_argument('--out', default=str(DEFAULT_OUT))
    args = parser.parse_args(argv)
    dataset = build(args.source)
    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(dataset, ensure_ascii=False, indent=1), encoding='utf-8', newline='\n')
    print('months %s, %d KPI actuals, %d irregularities -> %s' % (
        dataset['months_with_revenue'], len(dataset['kpi_actuals']), len(dataset['irregularities']), out))


if __name__ == '__main__':
    main()
