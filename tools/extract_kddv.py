# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""Turn the VTV Sales & Services planning files into one normalised dataset.

The customer handed over three documents (staff list, the Q3/2026 OKR decision
and the monthly KPI assignment workbook). The production data is entered from
them through the browser, and an independent acceptance spec checks the result
against them. Both need the same reading of the files, so the reading happens
once, here, and is written to JSON.

Nothing in this script talks to Odoo. It only reads spreadsheets.

Run:  python tools/extract_kddv.py [--source Docs/OKR] [--out uat/data/kddv_q3_2026.json]
"""
import argparse
import json
import pathlib
import re

_REPO = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = _REPO / 'Docs' / 'OKR'
DEFAULT_OUT = _REPO / 'uat' / 'data' / 'kddv_q3_2026.json'

STAFF_FILE = 'TT Nhân viên (hr.employee).xlsx'
KPI_FILE = 'Giao_KPI_Thang_T7-T8-T9.2026_Phong_KD.xlsx'
OKR_FILE = 'Giao nhiệm vụ OKR Quý III.2026.pdf'

COMPANY_NAME = 'Trung tâm Nền tảng và Dịch vụ số'
DEPARTMENT_CODE = 'KDDV'
MONTHS = (7, 8, 9)
MONTH_COLUMNS = {7: 5, 8: 6, 9: 7}
YEAR = 2026

# The KPI workbook lists each position's holders on its overview sheet. Two
# people are not covered by any assignment sheet and need an explicit rule.
NO_KPI_POSITION = 'LX'
LATE_JOINER = {
    # Staff note: "Chuyển về từ 1/9 nên tháng 8 không đánh giá". The staff
    # title is "Biên tập viên hạng III", which matches the BTV sheet.
    'Đinh Duy Phương': {'position': 'BTV', 'months': [9]},
}

# Official wording from Appendix 5 of the decision. The workbook repeats the
# structure in shortened form; the decision is the document that was signed.
OBJECTIVES = [
    {
        'code': 'O1', 'weight': 50.0,
        'name': 'Hoàn thành kế hoạch doanh thu Quý III/2026',
        'key_results': [
            {'code': 'O1.KR1', 'weight': 35.0, 'deadline': '2026-09-30',
             'name': 'Đảm bảo doanh thu Quý III/2026 đạt theo đúng kế hoạch.',
             'criterion': 'Doanh thu Quý III/2026: ≥150,56 tỷ; gồm 5 mảng: Tiếp phát sóng kênh; '
                          'Thương mại điện tử (VTVshop MG); DV TNND; FAST Channel; DV số.',
             'metric_type': 'number', 'unit': 'tỷ VNĐ', 'baseline': 0.0, 'target': 150.56},
            {'code': 'O1.KR2', 'weight': 10.0, 'deadline': '2026-09-07',
             'name': 'Ra mắt thành công VTVshop B2C: hạ tầng kỹ thuật hoàn thiện, thực hiện thành công '
                     'giao dịch TMĐT đầu tiên trên VTVgo.',
             'criterion': 'Tiếp cận tối thiểu 10 đối tác/đơn vị bán hàng quan tâm sàn TMĐT; thực hiện '
                          'thành công tối thiểu 01 GD TMĐT trong Quý III.',
             'metric_type': 'milestone',
             'milestones': ['Hạ tầng kỹ thuật VTVshop B2C hoàn thiện',
                            'Tiếp cận tối thiểu 10 đối tác/đơn vị bán hàng',
                            'VTVshop B2C LIVE ngày 07/09/2026',
                            'Thực hiện thành công tối thiểu 01 giao dịch TMĐT']},
            {'code': 'O1.KR3', 'weight': 5.0, 'deadline': '2026-09-30',
             'name': 'Có giải pháp và triển khai thu hồi các công nợ quá hạn, nợ xấu',
             'criterion': 'Hoàn thành việc triển khai thu hồi công nợ với 4 đối tác nợ quá hạn, nợ xấu '
                          'trong Quý III',
             'metric_type': 'number', 'unit': 'đối tác', 'baseline': 0.0, 'target': 4.0},
        ],
    },
    {
        'code': 'O2', 'weight': 20.0,
        'name': 'Bứt phá tăng trưởng người dùng trên VTVgo, khẳng định VTVgo là nền tảng số hàng đầu '
                'Việt Nam',
        'key_results': [
            {'code': 'O2.KR1', 'weight': 10.0, 'deadline': '2026-09-30',
             'name': 'Đảm bảo tăng số lượng người dùng hàng tháng (MAU) trên VTVGo',
             'criterion': 'Tăng từ ≥10M (cuối Quý 2) lên ≥15M (cuối Quý 3), được dẫn dắt bởi lượng người '
                          'dùng từ WC→VTVgoPlus và người dùng mới từ VTVshop B2C ra mắt vào 7/9',
             'metric_type': 'number', 'unit': 'triệu người dùng', 'baseline': 10.0, 'target': 15.0},
            {'code': 'O2.KR2', 'weight': 6.0, 'deadline': '2026-09-30',
             'name': 'Thực hiện chiến dịch tiếp thị, kết nối lại khách hàng (re-engagement) từ giai đoạn '
                     'hậu World Cup',
             'criterion': 'Triển khai trong vòng 7 ngày sau khi WC kết thúc; WC→VTVgoPlus ≥15%; Doanh thu '
                          'bình quân trên một người dùng (ARPU) toàn nền tảng ≥3.000đ/tháng',
             'metric_type': 'milestone',
             'milestones': ['Triển khai chiến dịch trong vòng 7 ngày sau khi WC kết thúc',
                            'Tỷ lệ chuyển đổi WC→VTVgoPlus ≥15%',
                            'ARPU toàn nền tảng ≥3.000đ/tháng']},
            {'code': 'O2.KR3', 'weight': 4.0, 'deadline': '2026-09-30',
             'name': 'Phát triển lượng khách hàng thân thiết trên VTVGo (VTVgo Loyalty); tích hợp Loyalty '
                     'với VTVshop B2C kích hoạt lưu lượng người dùng mới.',
             'criterion': 'Loyalty active +20% so với Quý 2; tỷ lệ đổi điểm ≥30%; tích hợp Loyalty với '
                          'VTVshop B2C',
             'metric_type': 'milestone',
             'milestones': ['Loyalty active +20% so với Quý 2',
                            'Tỷ lệ đổi điểm ≥30%',
                            'Tích hợp Loyalty với VTVshop B2C']},
        ],
    },
    {
        'code': 'O3', 'weight': 20.0,
        'name': 'Xây dựng và hoàn thiện mô hình kinh doanh mới trên VTVgo.',
        'key_results': [
            {'code': 'O3.KR1', 'weight': 8.0, 'deadline': '2026-09-30',
             'name': 'VTVshop B2C vận hành thành công từ 7/9; Quy trình vận hành được tài liệu hóa đầy đủ; '
                     'Hoàn thành lộ trình phát triển VTVshop B2C cho 2027',
             'criterion': 'Quy trình vận hành; Lộ trình phát triển được phê duyệt',
             'metric_type': 'milestone',
             'milestones': ['VTVshop B2C vận hành thành công từ 7/9',
                            'Quy trình vận hành được tài liệu hóa đầy đủ',
                            'Lộ trình phát triển VTVshop B2C 2027 được phê duyệt']},
            {'code': 'O3.KR2', 'weight': 7.0, 'deadline': '2026-09-30',
             'name': 'Thiết kế và thử nghiệm DV trải nghiệm nội dung trả phí mới (ngoài VTVgo Plus); Đảm '
                     'bảo có hợp đồng nguyên tắc (LOI) ký với đối tác Tài chính số (Finance) để triển khai '
                     'DV tài chính số trên VTVgo từ Q4/2026',
             'criterion': '≥1 gói DV mới thử nghiệm; ≥1 LOI tài chính số',
             'metric_type': 'milestone',
             'milestones': ['≥1 gói DV trải nghiệm nội dung trả phí mới thử nghiệm',
                            '≥1 LOI tài chính số được ký']},
            {'code': 'O3.KR3', 'weight': 5.0, 'deadline': '2026-09-30',
             'name': 'Mở rộng FAST Channel & Chuyên trang mới',
             'criterion': 'Mở rộng tối thiểu 01 chuyên trang mới',
             'metric_type': 'number', 'unit': 'chuyên trang', 'baseline': 0.0, 'target': 1.0},
        ],
    },
    {
        'code': 'O4', 'weight': 10.0,
        'name': 'Xây dựng và hoàn thiện các quy trình liên quan đến kinh doanh và dịch vụ',
        'key_results': [
            {'code': 'O4.KR1', 'weight': 10.0, 'deadline': '2026-09-30',
             'name': 'Hoàn thiện quy trình kinh doanh, cơ chế phối hợp với các đơn vị bên ngoài để đạt '
                     'được các mục tiêu về doanh thu',
             'criterion': 'Các quy trình kinh doanh và cơ chế phối hợp được phê duyệt và áp dụng',
             'metric_type': 'milestone',
             'milestones': ['Các quy trình kinh doanh được phê duyệt',
                            'Cơ chế phối hợp với đơn vị bên ngoài được phê duyệt',
                            'Các quy trình được áp dụng trong Quý III']},
        ],
    },
]

_KR_REF = re.compile(r'O(\d)\s*-\s*KR(\d)')
_GROUP_WEIGHT = re.compile(r'trọng số nhóm[^:]*:\s*(\d+(?:[.,]\d+)?)\s*%')
_LEADING_NUMBER = re.compile(r'^\s*(≥|≤|\+)?\s*(\d[\d.,]*)')
_NOT_APPLICABLE = {'—', '-', '–'}


def parse_number(text):
    """Vietnamese number formatting: '.' groups thousands, ',' marks decimals.

    '2.600' -> 2600, '15.000' -> 15000, '49,15' -> 49.15, '0,20' -> 0.2.
    A dot is a thousands separator only when exactly three digits follow it.
    """
    raw = text.strip().rstrip('.,')
    raw = re.sub(r'\.(?=\d{3}(?:\D|$))', '', raw)
    return float(raw.replace(',', '.'))


def parse_month_target(text):
    """Read one monthly target cell.

    Returns None when the KPI does not apply that month ('—'), otherwise a dict
    with the scoring direction, a numeric target and the verbatim wording.
    The verbatim text is always kept: '≥ 21,92 tỷ (gốc 20,92 + 1,00 bổ sung)'
    scores against 21.92, but the explanation is part of what was assigned.

    A cell that opens with a comparison or a number is numeric. Anything else
    is a milestone ('Dự thảo 3 quy trình') and is scored pass/fail.
    """
    if text is None:
        return None
    text = str(text).strip()
    if not text or text in _NOT_APPLICABLE:
        return None
    match = _LEADING_NUMBER.match(text)
    if not match:
        return {'direction': 'boolean', 'target': 1.0, 'note': text}
    sign, number = match.groups()
    value = parse_number(number)
    if sign == '≤':
        direction = 'lower'
    elif not sign and value == 0.0:
        # "0 vụ vi phạm": zero tolerance, lower is better with a zero target.
        direction = 'lower'
    else:
        direction = 'higher'
    return {'direction': direction, 'target': value, 'note': text}


def parse_weight(value):
    if isinstance(value, (int, float)):
        return float(value) * (100.0 if value <= 1 else 1.0)
    return parse_number(str(value).replace('%', ''))


def parse_kr_refs(text):
    """'O1-KR1 (MG) + O1-KR2 (B2C 10%)' -> ['O1.KR1', 'O1.KR2'], order kept, no repeats."""
    refs = []
    for objective, key_result in _KR_REF.findall(str(text or '')):
        code = 'O%s.KR%s' % (objective, key_result)
        if code not in refs:
            refs.append(code)
    return refs


def _cell(sheet, row, column):
    value = sheet.cell(row, column).value
    return value.strip() if isinstance(value, str) else value


def read_staff(path):
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=True)
    staff_sheet, department_sheet = book['Staff'], book['Department']
    employees = []
    for row in range(2, staff_sheet.max_row + 1):
        code = _cell(staff_sheet, row, 1)
        if not code:
            continue
        employees.append({
            'code': code,
            'name': _cell(staff_sheet, row, 2),
            'email': _cell(staff_sheet, row, 3),
            'department': _cell(staff_sheet, row, 4),
            'job_title': _cell(staff_sheet, row, 5),
            'manager': _cell(staff_sheet, row, 6),
            'note': _cell(staff_sheet, row, 7),
        })
    departments = []
    for row in range(2, department_sheet.max_row + 1):
        code = _cell(department_sheet, row, 1)
        if not code:
            continue
        departments.append({
            'code': code,
            'name': _cell(department_sheet, row, 2),
            'manager': _cell(department_sheet, row, 3),
            'headcount': _cell(department_sheet, row, 4),
            'parent': _cell(department_sheet, row, 5),
        })
    return employees, departments


def read_positions(book):
    """Section V of the overview sheet: position code, title, unit, holders."""
    sheet = book.worksheets[0]
    positions, in_section = [], False
    for row in range(1, sheet.max_row + 1):
        first = _cell(sheet, row, 1)
        if isinstance(first, str) and first.startswith('Mã VTVL'):
            in_section = True
            continue
        if not in_section or not first:
            continue
        holders = [name.strip() for name in str(_cell(sheet, row, 5) or '').split(',') if name.strip()]
        positions.append({
            'code': first,
            'title': _cell(sheet, row, 2),
            'unit': _cell(sheet, row, 3),
            'holders': holders,
            'has_kpi': first != NO_KPI_POSITION,
        })
    return positions


def read_sheet(sheet):
    """One position's assignment sheet -> KPI groups with their lines."""
    groups = []
    for row in range(6, sheet.max_row + 1):
        code = _cell(sheet, row, 1)
        if not isinstance(code, str):
            continue
        if code.startswith('B.'):
            weight = _GROUP_WEIGHT.search(code)
            label = code.split('(')[0].strip()
            groups.append({
                'code': label.split('—')[0].strip(),
                'name': label.split('—', 1)[1].strip() if '—' in label else label,
                'weight': parse_number(weight.group(1)),
                'lines': [],
            })
            continue
        if not re.match(r'^B\d\.\d+$', code) or not groups:
            continue
        link = _cell(sheet, row, 11)
        groups[-1]['lines'].append({
            'code': code,
            'name': _cell(sheet, row, 2),
            'unit': _cell(sheet, row, 3),
            'weight_in_group': parse_weight(_cell(sheet, row, 4)),
            'months': {str(month): parse_month_target(_cell(sheet, row, column))
                       for month, column in MONTH_COLUMNS.items()},
            'quarter_target': _cell(sheet, row, 8),
            'basis': _cell(sheet, row, 9),
            'measurement': _cell(sheet, row, 10),
            'okr_link': link,
            'kr_codes': parse_kr_refs(link),
        })
    return groups


def month_lines(group, month):
    """The lines of a group that apply in a month, with in-group weights rescaled.

    A line marked '—' for a month is not assigned that month (BTV B1.3 starts
    in September). The remaining lines keep their proportions and are scaled
    back to 100, so the group still carries its full share of the score. The
    original weight stays on the line for the record.
    """
    active = [line for line in group['lines'] if line['months'][str(month)]]
    total = sum(line['weight_in_group'] for line in active)
    result = []
    for line in active:
        weight = line['weight_in_group'] if abs(total - 100.0) < 0.001 \
            else round(line['weight_in_group'] * 100.0 / total, 2)
        result.append(dict(line, month_weight_in_group=weight, target=line['months'][str(month)]))
    return result


def build(source=DEFAULT_SOURCE):
    import openpyxl
    source = pathlib.Path(source)
    employees, departments = read_staff(source / STAFF_FILE)
    book = openpyxl.load_workbook(source / KPI_FILE, data_only=True)
    positions = read_positions(book)
    sheets = {sheet.title: read_sheet(sheet) for sheet in book.worksheets[1:]}

    holder_position = {name: position['code'] for position in positions for name in position['holders']}
    by_code = {position['code']: position for position in positions}
    head = next(department for department in departments if department['code'] == DEPARTMENT_CODE)['manager']

    for employee in employees:
        rule = LATE_JOINER.get(employee['name'])
        position_code = rule['position'] if rule else holder_position.get(employee['name'])
        employee['position'] = position_code
        employee['position_title'] = by_code[position_code]['title'] if position_code else None
        employee['kpi_months'] = [] if position_code in (None, NO_KPI_POSITION) else \
            (rule['months'] if rule else list(MONTHS))
        employee['manager'] = employee['manager'] or (None if employee['name'] == head else head)
        employee['performance_role'] = 'manager' if position_code in ('TP-KD', 'PPT-KD', 'PPT-MKT') else 'user'

    scorecards = []
    for month in MONTHS:
        for employee in employees:
            if month not in employee['kpi_months']:
                continue
            groups = [dict(group, lines=month_lines(group, month)) for group in sheets[employee['position']]]
            scorecards.append({'month': month, 'employee': employee['name'],
                               'position': employee['position'], 'groups': groups})

    return {
        'company': COMPANY_NAME,
        'department_code': DEPARTMENT_CODE,
        'year': YEAR,
        'quarter_weights': {'Q1': 20.0, 'Q2': 20.0, 'Q3': 25.0, 'Q4': 35.0},
        'departments': departments,
        'employees': employees,
        'positions': positions,
        'objectives': OBJECTIVES,
        'kpi_sheets': sheets,
        'scorecards': scorecards,
        'source_notes': [
            'Phần bổ sung +3,08 tỷ cho mảng Tiếp phát sóng kênh (Telco/ISP) mới được TP-KD xác nhận qua '
            'trao đổi ngày 08/9/2026, chưa có văn bản điều chỉnh chính thức.',
            'DV trải nghiệm nội dung: 3 tháng cộng 17,60 tỷ nhưng cột tổng quý ghi 17,55 tỷ; dữ liệu dùng '
            'số theo tháng.',
            'Đinh Duy Phương chuyển về từ 1/9: chỉ giao KPI tháng 9, theo phiếu Biên tập viên (BTV).',
            'BTV B1.3 (nội dung VTVshop B2C) chỉ áp dụng từ tháng 9; tháng 7–8 trọng số trong nhóm của '
            'các dòng còn lại được quy đổi theo tỷ lệ về 100%.',
        ],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--source', default=str(DEFAULT_SOURCE))
    parser.add_argument('--out', default=str(DEFAULT_OUT))
    options = parser.parse_args()
    data = build(options.source)
    out = pathlib.Path(options.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print('%s: %d employees, %d positions, %d scorecards'
          % (out, len(data['employees']), len(data['positions']), len(data['scorecards'])))


if __name__ == '__main__':
    main()
