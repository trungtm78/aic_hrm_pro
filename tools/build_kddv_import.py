# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""Write one import workbook per month from the extracted Sales & Services plan.

The workbooks use the product's own spreadsheet format (OKR_2026 /
KPI_CHI_TIET / PHAN_CONG) so the plan enters production through
Performance > Plan > Import Spreadsheet, the screen a customer would use.
The OKR sheet stays empty: objectives and key results are entered on their
forms, and KPIs refer to them by code.

Output folder is git-ignored (it names real people).

Run:  python tools/build_kddv_import.py [--dataset uat/data/kddv_q3_2026.json] [--out uat/data/import]
"""
import argparse
import json
import pathlib

_REPO = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_DATASET = _REPO / 'uat' / 'data' / 'kddv_q3_2026.json'
DEFAULT_OUT = _REPO / 'uat' / 'data' / 'import'

# Words the import term table already maps (aic_okr_kpi/data/aic_hrm_import_terms.xml).
DIRECTION_TERM = {'higher': 'Càng cao càng tốt', 'lower': 'Càng thấp càng tốt', 'boolean': 'Đạt/Không đạt'}
AGGREGATION_TERM = 'Cuối kỳ'

KPI_HEADER = ['KPI ID', 'Nhóm KPI', 'Objective', 'KPI', 'Chiều đo', 'Cách tổng hợp', 'Đơn vị',
              'Trọng số', 'Chỉ tiêu', 'Chủ trì', 'Nguồn đo', 'Ghi chú',
              'Trọng số nhóm (%)', 'Trọng số trong nhóm (%)', 'Mã KR', 'Chỉ tiêu (nguyên văn)']
OKR_HEADER = ['Mã Objective', 'Objective', 'Tỷ trọng O (%)', 'Mã KR', 'Key Result', 'Thước đo',
              'Target', 'Đơn vị', 'Chủ trì', 'Quý trọng tâm', 'Ưu tiên', 'Ghi chú']
ASSIGN_HEADER = ['TT', 'Họ và tên', 'Vị trí', 'Trách nhiệm trọng tâm', 'Objective liên quan',
                 'KPI giao', 'Tổng trọng số (%)']


def kpi_code(position, line_code):
    return 'KDDV.%s.%s' % (position, line_code)


def group_label(group):
    return '%s — %s' % (group['code'], group['name'])


def month_rows(dataset, month):
    """(kpi_rows, assign_rows) for one month, from the per-person scorecards."""
    titles = {position['code']: position['title'] for position in dataset['positions']}
    by_position = {}
    for card in dataset['scorecards']:
        if card['month'] == month:
            by_position.setdefault(card['position'], {'groups': card['groups'], 'people': []})
            by_position[card['position']]['people'].append(card['employee'])

    kpi_rows, assign_rows = [], []
    for position, entry in by_position.items():
        codes = []
        for group in entry['groups']:
            for line in group['lines']:
                target = line['target']
                code = kpi_code(position, line['code'])
                codes.append(code)
                share = line['month_weight_in_group']
                kr_codes = line['kr_codes']
                note = 'Căn cứ số liệu: %s. Gắn OKR: %s' % (line['basis'] or '—', line['okr_link'] or '—')
                kpi_rows.append([
                    code, group_label(group), '', line['name'], DIRECTION_TERM[target['direction']],
                    AGGREGATION_TERM, line['unit'], round(group['weight'] * share / 100.0, 4),
                    target['target'], entry['people'][0], line['measurement'], note,
                    group['weight'], share, kr_codes[0] if kr_codes else '', target['note'],
                ])
        assign_rows.append([
            len(assign_rows) + 1, ', '.join(entry['people']), titles[position], '', '',
            ','.join(codes), 100,
        ])
    return kpi_rows, assign_rows


def write_month(path, kpi_rows, assign_rows):
    import openpyxl
    book = openpyxl.Workbook()
    book.active.title = 'OKR_2026'
    book.active.append(OKR_HEADER)
    kpi = book.create_sheet('KPI_CHI_TIET')
    kpi.append(KPI_HEADER)
    for row in kpi_rows:
        kpi.append(row)
    assign = book.create_sheet('PHAN_CONG')
    assign.append(ASSIGN_HEADER)
    for row in assign_rows:
        assign.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--dataset', default=str(DEFAULT_DATASET))
    parser.add_argument('--out', default=str(DEFAULT_OUT))
    options = parser.parse_args()
    dataset = json.loads(pathlib.Path(options.dataset).read_text(encoding='utf-8'))
    for month in (7, 8, 9):
        kpi_rows, assign_rows = month_rows(dataset, month)
        path = pathlib.Path(options.out) / ('KDDV_KPI_T%d_%d.xlsx' % (month, dataset['year']))
        write_month(path, kpi_rows, assign_rows)
        print('%s: %d KPI rows, %d assignment rows' % (path, len(kpi_rows), len(assign_rows)))


if __name__ == '__main__':
    main()
