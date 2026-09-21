# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""Add the screen walkthrough to a progress deck somebody is already editing.

The report is rewritten by hand after it is generated, so this tool does not
rebuild it: it opens the file as it is and appends slides, in the same visual
style, showing each screen of the system with a short note on what it is for.
The point is that a reader who never opens the system can still see what the
system is made of.

Style is measured from a slide already in the file (title, lead, note column,
page number, footnote), so a deck restyled by hand keeps its look. Text is
written plainly: what the screen shows, what it is for, where it sits in the
menu. No adjectives, no flourish.

Run:  python tools/append_kddv_screens.py --deck "Docs/OKR/Reports/<tệp>.pptx"
"""
import argparse
import pathlib
import re

_REPO = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_IMAGES = _REPO / 'Docs' / 'OKR' / 'img_kddv'

# Measured from the deck this was written for; `probe_style` overrides each of
# these from the file itself when it recognises the matching shape.
STYLE = {
    'title': {'box': (0.58, 0.40, 12.17, 0.96), 'size': 30.0, 'bold': True, 'colour': '16334B'},
    'lead': {'box': (0.60, 1.40, 12.03, 0.67), 'size': 17.25, 'bold': False, 'colour': '637382'},
    'note_head': {'box': (8.43, 2.24, 4.19, 0.35), 'size': 20.25, 'bold': True, 'colour': '146B94'},
    'note_body': {'box': (8.43, 2.64, 4.19, 0.89), 'size': 17.25, 'bold': False, 'colour': '16334B'},
    'footnote': {'box': (0.60, 6.55, 12.03, 0.45), 'size': 12.75, 'bold': False, 'colour': '637382'},
    'page': {'box': (12.19, 7.06, 0.62, 0.25), 'size': 11.25, 'bold': False, 'colour': '637382'},
}
NOTE_ROWS = (2.24, 3.52, 4.80)      # top of each note in the right-hand column
PICTURE_BOX = (0.60, 2.05, 7.55, 4.40)

# The walkthrough, in the order the work happens. Each entry: the screenshot,
# the screen's name as the menu shows it, one line on what it answers, up to
# three notes, and where to find it. A note states a figure that is on the
# picture or a rule the reader can check - never what the screen is not.
# What stands behind each screen. A screen whose model holds nothing is a
# blank page in a report to the customer's leadership, so `check_data`
# refuses to include it rather than letting it ship empty.
BACKING = {
    '21-cycles': 'aic.hrm.cycle',
    '24-objective-with-key-results': 'aic.hrm.objective',
    '23-key-results': 'aic.hrm.key.result',
    '20-alignment-tree': 'aic.hrm.objective',
    '22-kpi-library': 'aic.hrm.kpi',
    '03-kpi-target-form': 'aic.hrm.kpi.target',
    '02-scorecard-form': 'aic.hrm.kpi.assignment',
    '08-invoice-list': 'account.move',
    '10-cost-entry': 'account.move',
    '11-metric-sources': 'aic.hrm.metric.source',
    '05-period-results': 'aic.hrm.kpi.period.result',
    '07-checkins': 'aic.hrm.checkin',
    '17-cockpit-month': 'aic.hrm.kpi.assignment',
    '12-department-report': 'aic.hrm.kpi.assignment',
    '18-executive-overview': 'aic.hrm.kpi.target',
    '19-progress-vs-plan': 'aic.hrm.kpi.target',
    '14-review-cycle': 'aic.hrm.review',
    '29-objective-contribution': 'aic.hrm.objective.contribution',
}

SCREENS = [
    ('21-cycles', 'Chu kỳ làm việc',
     'Mọi việc giao và mọi điểm số đều gắn vào một chu kỳ: năm, quý hoặc tháng.',
     [('Cây chu kỳ của đơn vị', 'Năm 2026 chứa Quý III/2026, quý chứa tháng 7, 8 và 9.'),
      ('Điểm cộng lên', 'Điểm tháng cộng lên quý theo trọng số, không nhập lại.'),
      ('Khoá kỳ khi chốt', 'Kỳ đã khoá thì sửa số phải qua điều chỉnh có phê duyệt.')],
     'Hiệu suất › Kế hoạch › Chu kỳ'),
    ('24-objective-with-key-results', 'Mục tiêu quý của phòng',
     'Bốn mục tiêu Quý III/2026 của phòng, mỗi mục tiêu có trọng số và người chịu trách nhiệm.',
     [('Trọng số quyết định', 'O1 trọng số 50%, O2 và O3 mỗi cái 20%, O4 10%.'),
      ('Đo bằng kết quả then chốt', 'O1 có 3 kết quả then chốt, trong đó doanh thu quý đã đạt 209,77 tỷ.'),
      ('Điểm tự tính', 'O1 đạt 70% nên đóng góp 35% vào điểm chung của quý.')],
     'Hiệu suất › Kế hoạch › Mục tiêu'),
    ('23-key-results', 'Kết quả then chốt',
     'Mười kết quả then chốt của Quý III, đo bằng con số, bằng mốc công việc hoặc bằng xong và chưa xong.',
     [('Hiện mới 1/10 có số', 'Chỉ kết quả doanh thu đã được báo cáo, chín cái còn lại chưa.'),
      ('Ghi rõ mốc và đích', 'Ví dụ MAU: xuất phát 10 triệu, đích 15 triệu người dùng.'),
      ('Chưa báo là chưa chấm', 'Kết quả chưa ai báo ghi "chưa chấm điểm", không tính thành 0.')],
     'Hiệu suất › Kế hoạch › Kết quả then chốt'),
    ('20-alignment-tree', 'Cây liên kết mục tiêu',
     'Xem mục tiêu của phòng nối xuống kết quả then chốt và xuống chỉ tiêu của từng người.',
     [('Nhìn hết trong một trang', 'Thấy mục tiêu nào đã có người nhận, mục tiêu nào chưa.'),
      ('Dùng khi giao việc', 'Rà xem 418 chỉ tiêu đã phủ hết bốn mục tiêu quý chưa.'),
      ('Đổi chu kỳ để so', 'Chọn quý hoặc tháng khác để xem cây của kỳ đó.')],
     'Hiệu suất › Kế hoạch › Cây liên kết mục tiêu'),
    ('22-kpi-library', 'Thư viện KPI',
     'Kho 204 chỉ tiêu dùng lại cho kỳ sau, trong đó 90 chỉ tiêu riêng của phòng.',
     [('Không gõ lại mỗi kỳ', 'Chọn từ thư viện rồi gán cho người và cho kỳ.'),
      ('Có sẵn cách đo', 'Mỗi chỉ tiêu ghi đơn vị và chiều tốt lên hay giảm xuống.'),
      ('Dùng chung được', 'Phòng khác triển khai sau lấy lại chính các chỉ tiêu này.')],
     'Hiệu suất › Cấu hình › Thư viện KPI'),
    ('03-kpi-target-form', 'Chỉ tiêu KPI của một người',
     'Một dòng giao việc: giao cho ai, kỳ nào, bao nhiêu, và lấy số thực hiện ở đâu.',
     [('Chỉ tiêu và đơn vị', 'Ví dụ doanh thu Telco tháng 7: 21,92 tỷ đồng.'),
      ('Nguồn số liệu', 'Dòng này gắn nguồn "DT Telco/ISP (TK 51131)" nên số tự về.'),
      ('Kết quả từng kỳ', 'Tab bên dưới giữ số thực hiện của từng tháng và trạng thái xác nhận.')],
     'Hiệu suất › Kế hoạch › Chỉ tiêu KPI'),
    ('02-scorecard-form', 'Phiếu giao KPI của một người',
     'Các chỉ tiêu của một người trong một tháng, tổng trọng số 100.',
     [('Xếp theo nhóm', 'Nhóm doanh thu và nhóm quản trị, mỗi nhóm một phần trọng số.'),
      ('Trọng số quyết định điểm', 'Dòng doanh thu tổng của trưởng phòng chiếm 32 trên 100.'),
      ('Điểm kèm độ phủ', 'Phiếu tháng 7 của trưởng phòng: 80% trọng số có số, điểm 100% trên phần đó.')],
     'Hiệu suất › Kế hoạch › Bảng điểm cá nhân'),
    ('08-invoice-list', 'Hoá đơn bán hàng',
     'Doanh thu vào hệ thống bằng hoá đơn như nghiệp vụ kế toán, không nhập thẳng vào KPI.',
     [('96 hoá đơn đã vào sổ', 'Từ tháng 1 đến tháng 8, của 21 đối tác.'),
      ('Tách theo mảng', 'Mỗi dòng hoá đơn vào đúng tài khoản doanh thu của mảng.'),
      ('Truy ngược được', 'Từ điểm KPI mở ra đúng hoá đơn đã sinh ra con số đó.')],
     'Hoá đơn › Khách hàng › Hoá đơn'),
    ('10-cost-entry', 'Bút toán chi phí tháng',
     'Chi phí từng tháng vào sổ theo khoản mục, dùng cho chỉ tiêu chi phí của phòng.',
     [('7 bút toán, 104,4 tỷ', 'Mỗi tháng từ tháng 1 đến tháng 7 một bút toán tổng hợp.'),
      ('Tách theo khoản mục', 'Nhân công, sản xuất chung, tài chính tách riêng từng dòng.'),
      ('Nối vào KPI', 'Chỉ tiêu chi phí cấp phòng lấy số từ chính bút toán này.')],
     'Kế toán › Sổ nhật ký'),
    ('11-metric-sources', 'Nguồn số liệu tự động',
     'Tám nguồn đang chạy, lấy số từ sổ kế toán cho các chỉ tiêu doanh thu, chi phí và lợi nhuận.',
     [('Khai báo một lần', 'Ghi rõ tài khoản và cách cộng, ví dụ tài khoản 51131 cho Telco.'),
      ('Chạy theo từng kỳ', 'Đến kỳ nào hệ thống lấy đúng số phát sinh của kỳ đó.'),
      ('Phủ 21% trọng số', 'Phần này không ai gõ lại nên không lệch được với sổ.')],
     'Hiệu suất › Cấu hình › Nguồn số liệu'),
    ('05-period-results', 'Kết quả theo kỳ',
     'Số thực hiện của từng chỉ tiêu trong từng tháng, kèm trạng thái và người xác nhận.',
     [('38 số đã xác nhận', 'Chỉ số đã xác nhận mới được tính vào điểm.'),
      ('Nháp chưa vào điểm', 'Số mới nhập nằm ở trạng thái nháp cho tới khi quản lý duyệt.'),
      ('Sửa phải nêu lý do', 'Huỷ xác nhận bắt buộc ghi lý do và được lưu lại.')],
     'Hiệu suất › Thực hiện › Kết quả theo kỳ'),
    ('07-checkins', 'Check-in tiến độ',
     'Người phụ trách báo tiến độ của kết quả then chốt, kèm mức độ tin tưởng đạt đích.',
     [('Ghi người và thời điểm', 'Mỗi lần check-in lưu ai báo và báo lúc nào.'),
      ('Vào thẳng điểm', 'Giá trị mới cập nhật ngay kết quả then chốt và điểm mục tiêu.'),
      ('Thấy được bỏ quên', 'Kết quả lâu không ai báo bị đánh dấu để nhắc.')],
     'Hiệu suất › Thực hiện › Check-in'),
    ('17-cockpit-month', 'Bàn điều hành theo tháng',
     'Cùng màn hình bàn điều hành nhưng xem một tháng, có thêm khối phiếu giao KPI.',
     [('Tháng dùng mục tiêu quý', 'Tháng 7 không giao mục tiêu riêng nên lấy mục tiêu Quý III.'),
      ('Khối KPI của tháng', 'Tháng 7: 7/19 phiếu có số, điểm 83% trên phần đó, độ phủ 22%.'),
      ('Ghi rõ nguồn', 'Dòng dưới tiêu đề nói số liệu lấy từ chu kỳ nào.')],
     'Hiệu suất › Bàn điều hành'),
    ('12-department-report', 'Bảng điểm phòng ban',
     'Điểm KPI và điểm mục tiêu bình quân của phòng theo từng kỳ.',
     [('Số người đã có số liệu', 'Tháng 7 và tháng 8 đều là 7 trên 19 người.'),
      ('Hai loại điểm', 'Điểm tổng hợp 18% và điểm trên phần có số liệu 83% đặt cạnh nhau.'),
      ('Điểm mục tiêu lấy từ quý', 'Cột cuối ghi rõ 35% là điểm của Quý III/2026.')],
     'Hiệu suất › Báo cáo › Bảng điểm phòng ban'),
    ('18-executive-overview', 'Tổng quan lãnh đạo',
     'Kế hoạch và thực hiện theo từng tháng của kỳ, đọc trước cuộc họp.',
     [('Đặt cạnh nhau', 'Cột kế hoạch và cột thực hiện của cùng một tháng.'),
      ('Quý III', 'Kế hoạch 150,51 tỷ, đã thực hiện 209,77 tỷ trong hai tháng đầu quý.'),
      ('Mở ra chi tiết', 'Bấm vào một tháng để xem số liệu đằng sau.')],
     'Hiệu suất › Báo cáo › Tổng quan lãnh đạo'),
    ('19-progress-vs-plan', 'Tiến độ so với kế hoạch',
     'Tiến độ của từng chỉ tiêu, lọc theo kỳ, theo phòng hoặc theo từng người.',
     [('Ba cách xem', 'Dạng bảng, biểu đồ và bảng xoay trên cùng dữ liệu.'),
      ('Lọc theo nhu cầu', 'Chọn chu kỳ, phòng ban hoặc một người cụ thể.'),
      ('Xuất ra Excel', 'Lấy dữ liệu ra tệp khi cần gửi đi.')],
     'Hiệu suất › Báo cáo › Tiến độ so với kế hoạch'),
    ('29-objective-contribution', 'Đóng góp vào mục tiêu',
     'Mỗi mục tiêu của phòng do những ai gánh, và mỗi người đang đạt bao nhiêu trên phần đã có số liệu.',
     [('Đọc theo mục tiêu', 'Gom nhóm theo mục tiêu rồi tới từng người trong phòng.'),
      ('Trọng số cam kết', 'Ví dụ tháng 7: trưởng phòng dành 89/100 trọng số phiếu cho mục tiêu doanh thu.'),
      ('Không cộng dồn', 'Một chỉ tiêu giao cho nhiều người cùng gánh nên các dòng không cộng lại được.')],
     'Hiệu suất › Báo cáo › Đóng góp vào mục tiêu'),
    ('14-review-cycle', 'Chu kỳ đánh giá',
     'Đợt đánh giá của một chu kỳ hiệu suất và toàn bộ phiếu sinh ra từ đó.',
     [('21 phiếu Quý III', 'Một thao tác sinh phiếu cho cả phòng.'),
      ('Bốn chặng', 'Tự đánh giá, quản lý đánh giá, hiệu chỉnh, chốt kết quả.'),
      ('Điểm KPI tự vào phiếu', 'Lấy từ phiếu giao KPI của quý, không chép tay.')],
     'Hiệu suất › Đánh giá › Chu kỳ đánh giá'),
]

MAP_SLIDE = {
    'title': 'Hệ thống gồm những màn hình nào',
    'lead': 'Bốn nhóm màn hình theo đúng thứ tự công việc trong một kỳ.',
    'blocks': [
        ('1. Lập kế hoạch',
         'Khai báo chu kỳ, giao mục tiêu quý cho phòng, giao chỉ tiêu KPI theo tháng cho từng người.'),
        ('2. Cập nhật kết quả',
         'Doanh thu và chi phí vào sổ kế toán; chỉ tiêu có nguồn thì lấy số tự động, còn lại người phụ trách nhập; quản lý xác nhận.'),
        ('3. Giám sát',
         'Bàn điều hành, bảng điểm phòng ban, tiến độ so với kế hoạch, cảnh báo và biên bản họp rà soát.'),
        ('4. Đánh giá',
         'Mở chu kỳ đánh giá, chạy các chặng, hiệu chỉnh giữa các nhóm rồi chốt kết quả và lập kế hoạch phát triển.'),
    ],
    'footnote': 'Các trang sau là ảnh chụp từng màn hình trên hệ thống của đơn vị.',
}

# Words that make a slide read like it was written by a machine rather than by
# somebody who did the work.
BANNED = ('mạnh mẽ', 'vượt trội', 'đột phá', 'toàn diện', 'tối ưu', 'nâng tầm',
          'chìa khoá', 'chìa khóa', 'bức tranh', 'đồng hành', 'giải pháp tổng thể',
          'đáng kể', 'vô cùng', 'hết sức', 'tuyệt vời', 'linh hoạt và', 'sâu sắc')


def check_wording(entries=SCREENS, extra=()):
    """Keep the walkthrough in plain words."""
    texts = list(extra)
    for _image, title, lead, notes, footnote in entries:
        texts += [title, lead, footnote] + [text for note in notes for text in note]
    bad = sorted({word for text in texts for word in BANNED if word in text.lower()})
    if bad:
        raise SystemExit('Chữ nghe như máy viết: %s' % ', '.join(bad))
    return len(texts)


def check_data(client, entries=None):
    """Refuse a screen that would go out as an empty page.

    Four screens - alert rules, review meetings, calibration and development
    plans - held no records at all. Sending those to a customer's leadership
    says the system is empty; the honest answer is to leave them out until
    the work behind them has actually happened.
    """
    entries = SCREENS if entries is None else entries
    empty = []
    for name, title, *_rest in entries:
        model = BACKING.get(name)
        if not model:
            empty.append('%s: chưa khai báo dữ liệu đứng sau' % name)
            continue
        if not client.call(model, 'search_count', [[]]):
            empty.append('%s ("%s"): %s chưa có bản ghi nào' % (name, title, model))
    if empty:
        raise SystemExit('Màn hình chưa có dữ liệu, không đưa vào báo cáo:\n  '
                         + '\n  '.join(empty))
    return len(entries)


def check_images(images, entries=SCREENS):
    missing = [name for name, *_rest in entries if not (images / f'{name}.png').exists()]
    if missing:
        raise SystemExit('Thiếu ảnh màn hình: %s.\nChạy "cd uat && node capture_kddv_guide.mjs" '
                         'rồi chạy lại.' % ', '.join(missing))


FONT_FILE = pathlib.Path('C:/Windows/Fonts/segoeui.ttf')
LINE_SPACING = 1.22          # PowerPoint's single spacing for this face


def wrap(text, size_pt, width_in, font_file=FONT_FILE):
    """Break `text` the way PowerPoint will, measured with the real font.

    Guessing characters-per-line is what put a four-line note in a box three
    lines tall, and the overflow landed on top of the next heading.
    """
    from PIL import ImageFont
    font = ImageFont.truetype(str(font_file), int(round(size_pt * 4)))  # 4x for precision
    limit = width_in * 72 * 4
    lines, current = [], ''
    for word in text.split():
        candidate = (current + ' ' + word).strip()
        if current and font.getlength(candidate) > limit:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def fits(text, size_pt, box, font_file=FONT_FILE):
    """Whether the text stays inside its box (width, height in inches)."""
    if not font_file.exists():
        return True, 0, 0
    lines = wrap(text, size_pt, box[0], font_file)
    needed = len(lines) * size_pt * LINE_SPACING / 72.0
    return needed <= box[1] + 0.02, len(lines), needed


def check_fit(entries=None, style=None):
    """Refuse to write a slide whose note would run over the next one."""
    entries = SCREENS if entries is None else entries
    style = style or STYLE
    head_box = (style['note_head']['box'][2], style['note_head']['box'][3])
    body_box = (style['note_body']['box'][2], style['note_body']['box'][3])
    title_box = (style['title']['box'][2], style['title']['box'][3])
    lead_box = (style['lead']['box'][2], style['lead']['box'][3])
    long_ones = []
    for name, title, lead, notes, where in entries:
        checks = [(title, style['title']['size'], title_box, 'tiêu đề'),
                  (lead, style['lead']['size'], lead_box, 'dòng dẫn'),
                  ('Vị trí trên menu: ' + where, style['footnote']['size'],
                   (style['footnote']['box'][2], style['footnote']['box'][3]), 'chân trang')]
        for head, body in notes:
            checks.append((head, style['note_head']['size'], head_box, 'tiêu đề ghi chú'))
            checks.append((body, style['note_body']['size'], body_box, 'ghi chú'))
        for text, size, box, what in checks:
            ok, lines, needed = fits(text, size, box)
            if not ok:
                long_ones.append('%s · %s: "%s" cần %.2f" (khung %.2f", %s dòng)'
                                 % (name, what, text[:48], needed, box[1], lines))
    if long_ones:
        raise SystemExit('Chữ dài hơn khung, sẽ tràn sang mục dưới:\n  '
                         + '\n  '.join(long_ones))
    return True


def retouch(deck_path, replacements, out=None, allow_overflow=False):
    """Rewrite named notes in a deck without touching its look.

    A note that only says what a number is not ("không phải tỷ lệ hoàn thành
    dự án") leaves the reader where they started. These replacements put the
    arithmetic in its place: where the figure comes from and what it counts.
    Text is written into the first run so the size, weight and colour the deck
    already uses are kept.
    """
    from pptx import Presentation
    deck = Presentation(str(deck_path))
    done, seen = [], set()
    for index, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            current = shape.text_frame.text.strip()
            if current not in replacements:
                continue
            # The same sentence can sit on several slides; a rewrite that
            # stopped at the first one would leave the deck saying two things.
            wanted = replacements[current]
            seen.add(current)
            paragraph = shape.text_frame.paragraphs[0]
            runs = paragraph.runs
            if not runs:
                continue
            # Measured against the box it is going into: a replacement one
            # line too long spills onto the heading below it, which is what
            # the customer saw on the quarter overview slide.
            from pptx.util import Emu
            size = runs[0].font.size.pt if runs[0].font.size else 18.0
            box = (Emu(shape.width).inches, Emu(shape.height).inches)
            room, lines, needed = fits(wanted, size, box)
            if not room and not allow_overflow:
                raise SystemExit(
                    'Trang %s: "%s" cần %.2f" (%s dòng) nhưng khung chỉ cao %.2f". '
                    'Hãy viết ngắn lại.' % (index, wanted[:60], needed, lines, box[1]))
            runs[0].text = wanted
            for extra in runs[1:]:
                extra.text = ''
            for other in shape.text_frame.paragraphs[1:]:
                for run in other.runs:
                    run.text = ''
            done.append((index, current, wanted))
    unknown = set(replacements) - seen
    if unknown:
        raise SystemExit('Không tìm thấy các đoạn chữ cần sửa: %s'
                         % ' | '.join(sorted(unknown)))
    deck.save(str(out or deck_path))
    return done


def set_font(deck_path, name='Segoe UI', out=None):
    """Put the whole deck on a font the reader's machine actually has.

    A deck carrying a font nobody has installed is re-laid out by PowerPoint
    when it opens: lines break in new places and boxes that fitted stop
    fitting. Segoe UI ships with Windows and covers Vietnamese, so the file
    looks on their screen the way it looks here. Sizes, weights and colours
    are left alone.
    """
    from pptx import Presentation
    deck = Presentation(str(deck_path))
    touched = 0

    def paint(frame):
        nonlocal touched
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = name
                touched += 1

    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                paint(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        paint(cell.text_frame)
            if shape.has_chart:
                shape.chart.font.name = name
    deck.save(str(out or deck_path))
    return touched


def replace_picture(deck_path, title, image, out=None):
    """Swap the picture on the slide with this title, keeping its frame.

    A screenshot is taken again whenever the screen behind it changes - the
    alignment tree was captured while it still opened on an empty month - and
    the slide around it must not move when the new one goes in.
    """
    from pptx import Presentation
    from pptx.util import Emu, Inches
    deck = Presentation(str(deck_path))
    for index, slide in enumerate(deck.slides, start=1):
        titles = [shape.text_frame.text for shape in slide.shapes
                  if shape.has_text_frame and Emu(shape.top).inches < 0.9]
        if title not in titles:
            continue
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        if not pictures:
            raise SystemExit('Trang %s ("%s") không có ảnh để thay.' % (index, title))
        old_picture = pictures[0]
        box = (Emu(old_picture.left).inches, Emu(old_picture.top).inches,
               Emu(old_picture.width).inches, Emu(old_picture.height).inches)
        # Keep the frame the slide was laid out around: same centre, same room.
        frame = (PICTURE_BOX[0], box[1], PICTURE_BOX[2], box[3])
        placed = _fit(image, frame)
        old_picture._element.getparent().remove(old_picture._element)
        slide.shapes.add_picture(str(image), Inches(placed[0]), Inches(placed[1]),
                                 width=Inches(placed[2]))
        deck.save(str(out or deck_path))
        return index
    raise SystemExit('Không tìm thấy trang có tiêu đề "%s".' % title)


def scan_overflow(deck_path):
    """Every text box in the deck that needs more room than it has.

    Run it after editing by hand: PowerPoint shows the overflow only when the
    slide is open, and printed or exported it lands on top of whatever sits
    below.
    """
    from pptx import Presentation
    from pptx.util import Emu
    found = []
    deck = Presentation(str(deck_path))
    for index, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            runs = [run for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs]
            if not runs:
                continue
            size = runs[0].font.size.pt if runs[0].font.size else 18.0
            box = (Emu(shape.width).inches, Emu(shape.height).inches)
            text = ' '.join(shape.text_frame.text.split())
            room, lines, needed = fits(text, size, box)
            if not room:
                found.append({'slide': index, 'text': text, 'lines': lines,
                              'needed': needed, 'height': box[1]})
    return found


def probe_style(slide):
    """Read the deck's own typography instead of imposing ours.

    A deck edited by hand keeps its look: we match each shape by where it sits
    and copy its size, weight and colour.
    """
    from pptx.util import Emu
    found = {}
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        left, top = Emu(shape.left).inches, Emu(shape.top).inches
        width = Emu(shape.width).inches
        runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        if not runs:
            continue
        font = runs[0].font
        role = None
        if top < 0.9 and width > 8:
            role = 'title'
        elif 1.1 < top < 1.8 and width > 8:
            role = 'lead'
        elif top > 6.3 and left < 2 and width > 8:
            role = 'footnote'
        elif left > 11.5 and top > 6.5:
            role = 'page'
        if role and role not in found:
            found[role] = {
                'box': (left, top, width, Emu(shape.height).inches),
                'size': font.size.pt if font.size else STYLE[role]['size'],
                'bold': bool(font.bold),
                'colour': str(font.color.rgb) if font.color and font.color.type is not None
                          else STYLE[role]['colour'],
                'name': font.name,
            }
    style = {role: dict(values) for role, values in STYLE.items()}
    for role, values in found.items():
        style[role].update(values)
    for role in ('note_head', 'note_body'):
        style[role].setdefault('name', found.get('lead', {}).get('name'))
    return style


def _text(slide, text, box, spec, Inches, Pt):
    from pptx.dml.color import RGBColor
    frame_box = slide.shapes.add_textbox(Inches(box[0]), Inches(box[1]),
                                         Inches(box[2]), Inches(box[3]))
    frame = frame_box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(spec['size'])
    run.font.bold = spec['bold']
    if spec.get('name'):
        run.font.name = spec['name']
    run.font.color.rgb = RGBColor.from_string(spec['colour'])
    return frame_box


def _fit(png, box):
    from PIL import Image
    with Image.open(png) as picture:
        width, height = picture.size
    scale = min(box[2] / width, box[3] / height)
    drawn = (width * scale, height * scale)
    return (box[0] + (box[2] - drawn[0]) / 2, box[1] + (box[3] - drawn[1]) / 2, *drawn)


def append(deck_path, images, entries=SCREENS, out=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    check_images(images, entries)
    check_fit(entries)
    check_wording(entries, extra=[MAP_SLIDE['title'], MAP_SLIDE['lead'],
                                  MAP_SLIDE['footnote']]
                  + [text for block in MAP_SLIDE['blocks'] for text in block])
    deck = Presentation(str(deck_path))
    reference = deck.slides[1] if len(deck.slides) > 1 else deck.slides[0]
    style = probe_style(reference)
    layout = reference.slide_layout
    number = len(deck.slides)

    def page(slide):
        nonlocal number
        number += 1
        _text(slide, '%02d' % number, style['page']['box'], style['page'], Inches, Pt)

    # The map first: what the following screens add up to.
    slide = deck.slides.add_slide(layout)
    _text(slide, MAP_SLIDE['title'], style['title']['box'], style['title'], Inches, Pt)
    _text(slide, MAP_SLIDE['lead'], style['lead']['box'], style['lead'], Inches, Pt)
    top = 2.19
    for headline, body in MAP_SLIDE['blocks']:
        _text(slide, headline, (0.62, top, 11.93, 0.35), style['note_head'], Inches, Pt)
        _text(slide, body, (0.62, top + 0.39, 11.93, 0.60), style['note_body'], Inches, Pt)
        top += 1.0
    _text(slide, MAP_SLIDE['footnote'], style['footnote']['box'], style['footnote'], Inches, Pt)
    page(slide)

    for name, title, lead, notes, where in entries:
        slide = deck.slides.add_slide(layout)
        _text(slide, title, style['title']['box'], style['title'], Inches, Pt)
        _text(slide, lead, style['lead']['box'], style['lead'], Inches, Pt)
        placed = _fit(images / f'{name}.png', PICTURE_BOX)
        slide.shapes.add_picture(str(images / f'{name}.png'), Inches(placed[0]),
                                 Inches(placed[1]), width=Inches(placed[2]))
        for (headline, body), row in zip(notes, NOTE_ROWS):
            head_box = (style['note_head']['box'][0], row, style['note_head']['box'][2], 0.35)
            body_box = (style['note_body']['box'][0], row + 0.40,
                        style['note_body']['box'][2], 0.89)
            _text(slide, headline, head_box, style['note_head'], Inches, Pt)
            _text(slide, body, body_box, style['note_body'], Inches, Pt)
        _text(slide, 'Vị trí trên menu: ' + where, style['footnote']['box'],
              style['footnote'], Inches, Pt)
        page(slide)

    target = pathlib.Path(out) if out else pathlib.Path(deck_path)
    deck.save(str(target))
    return {'added': len(entries) + 1, 'total': len(deck.slides), 'file': target}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--deck', required=True, help='tệp .pptx đang dùng')
    parser.add_argument('--images', default=str(DEFAULT_IMAGES))
    parser.add_argument('--out', default=None, help='ghi ra tệp khác, mặc định ghi đè')
    parser.add_argument('--check', action='store_true',
                        help='chỉ kiểm tra chữ có tràn khung không, không sửa gì')
    parser.add_argument('--font', default=None,
                        help='đổi phông cả tệp, ví dụ "Segoe UI"; mặc định giữ nguyên')
    args = parser.parse_args()
    if args.check:
        spills = scan_overflow(pathlib.Path(args.deck))
        for spill in spills:
            print('trang %s: cần %.2f" trong khung %.2f" (%s dòng): %s'
                  % (spill['slide'], spill['needed'], spill['height'],
                     spill['lines'], spill['text'][:70]))
        print('%s ô chữ tràn khung' % len(spills))
        return
    written = append(pathlib.Path(args.deck), pathlib.Path(args.images), out=args.out)
    if args.font:
        runs = set_font(written['file'], args.font)
        print('%s: đổi phông %s cho %s đoạn chữ' % (written['file'], args.font, runs))
    print('%s: thêm %s trang, tổng %s trang'
          % (written['file'], written['added'], written['total']))


if __name__ == '__main__':
    main()


def remove_slides(deck_path, titles, out=None):
    """Take slides out of a deck and close the gap in the page numbers.

    A screen with no records behind it is a blank page in a report to the
    customer's leadership: better absent than empty.
    """
    from pptx import Presentation
    from pptx.util import Emu
    deck = Presentation(str(deck_path))
    wanted = set(titles)
    dropped, keep = [], []
    slides = deck.slides._sldIdLst
    for slide, element in zip(list(deck.slides), list(slides)):
        heading = [shape.text_frame.text for shape in slide.shapes
                   if shape.has_text_frame and Emu(shape.top).inches < 0.9
                   and Emu(shape.width).inches > 8]
        if heading and heading[0] in wanted:
            dropped.append(heading[0])
            deck.part.drop_rel(element.rId)
            slides.remove(element)
        else:
            keep.append(slide)
    missing = wanted - set(dropped)
    if missing:
        raise SystemExit('Không tìm thấy trang: %s' % ', '.join(sorted(missing)))
    for number, slide in enumerate(keep, start=1):
        for shape in slide.shapes:
            if (shape.has_text_frame and Emu(shape.left).inches > 11.5
                    and shape.text_frame.text.strip().isdigit()):
                runs = shape.text_frame.paragraphs[0].runs
                if runs:
                    runs[0].text = '%02d' % number
                    for extra in runs[1:]:
                        extra.text = ''
    deck.save(str(out or deck_path))
    return dropped


TABLE_STYLE = {'head_fill': '16334B', 'head_text': 'FFFFFF', 'body_fill': 'EDF3F6',
               'body_text': '16334B', 'size': 13.5, 'head_size': 14.0}


def add_table_slide(deck_path, title, lead, headers, rows, widths, footnote='',
                    after=None, out=None, style=None):
    """One more slide in the deck's own shape: title, lead, a table, a note.

    Built for figures the report was missing rather than for a screenshot -
    per-person results, where a picture of a list would be unreadable at
    slide size.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    deck = Presentation(str(deck_path))
    reference = deck.slides[1] if len(deck.slides) > 1 else deck.slides[0]
    probed = style or probe_style(reference)
    table_style = dict(TABLE_STYLE)
    slide = deck.slides.add_slide(reference.slide_layout)
    _text(slide, title, probed['title']['box'], probed['title'], Inches, Pt)
    if lead:
        _text(slide, lead, probed['lead']['box'], probed['lead'], Inches, Pt)
    top = 2.10
    height = min(4.2, 0.42 * (len(rows) + 1))
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.60),
                                   Inches(top), Inches(12.12), Inches(height))
    table = shape.table
    for element in shape._element.graphic.graphicData.tbl.tblPr.findall('.//{*}tableStyleId'):
        element.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    total = sum(widths)
    for index, weight in enumerate(widths):
        table.columns[index].width = int(Inches(12.12).emu * weight / total)
    for column, header in enumerate(headers):
        _table_cell(table.cell(0, column), header, table_style['head_size'], True,
                    table_style['head_text'], table_style['head_fill'],
                    probed['lead'].get('name'), Pt, RGBColor, PP_ALIGN, MSO_ANCHOR,
                    'l' if column == 0 else 'c')
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            _table_cell(table.cell(row_index, column), str(value), table_style['size'], False,
                        table_style['body_text'], table_style['body_fill'],
                        probed['lead'].get('name'), Pt, RGBColor, PP_ALIGN, MSO_ANCHOR,
                        'l' if column == 0 else 'c')
    if footnote:
        _text(slide, footnote, probed['footnote']['box'], probed['footnote'], Inches, Pt)
    number = len(deck.slides) if after is None else after + 1
    _text(slide, '%02d' % number, probed['page']['box'], probed['page'], Inches, Pt)
    if after is not None:
        slides = deck.slides._sldIdLst
        element = slides[-1]
        slides.remove(element)
        slides.insert(after, element)
    deck.save(str(out or deck_path))
    return len(deck.slides)


def _table_cell(cell, text, size, bold, colour, fill, font, Pt, RGBColor,
                PP_ALIGN, MSO_ANCHOR, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Pt(8)
    cell.margin_top = cell.margin_bottom = Pt(4)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER,
                           'r': PP_ALIGN.RIGHT}[align]
    run = paragraph.add_run()
    run.text = text
    run.font.size, run.font.bold = Pt(size), bold
    if font:
        run.font.name = font
    run.font.color.rgb = RGBColor.from_string(colour)


def renumber(deck_path, out=None):
    """Put the page numbers back in order after slides move."""
    from pptx import Presentation
    from pptx.util import Emu
    deck = Presentation(str(deck_path))
    for number, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if (shape.has_text_frame and Emu(shape.left).inches > 11.5
                    and shape.text_frame.text.strip().isdigit()):
                runs = shape.text_frame.paragraphs[0].runs
                if runs:
                    runs[0].text = '%02d' % number
                    for extra in runs[1:]:
                        extra.text = ''
    deck.save(str(out or deck_path))
    return len(deck.slides)
