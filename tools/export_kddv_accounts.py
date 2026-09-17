# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""Write the account sheet handed to the customer.

Input is the account record the browser run keeps (uat/data/kddv_accounts.json)
plus the extracted dataset for titles. Only accounts that were proven by an
actual login are written: a sheet that lists a password nobody has tried is a
support call waiting to happen.

The output contains passwords. Its folder is git-ignored; never move it.

Run:  python tools/export_kddv_accounts.py
"""
import argparse
import json
import pathlib

_REPO = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_ACCOUNTS = _REPO / 'uat' / 'data' / 'kddv_accounts.json'
DEFAULT_DATASET = _REPO / 'uat' / 'data' / 'kddv_q3_2026.json'
DEFAULT_OUT = _REPO / 'Docs' / 'OKR' / 'Tai_khoan_KDDV_okr.aipower.vn.xlsx'
SITE = 'https://okr.aipower.vn'

HEADERS = ['STT', 'Mã NV', 'Họ và tên', 'Chức danh', 'Vị trí việc làm', 'Tên đăng nhập (email)',
           'Mật khẩu ban đầu', 'Quyền trên hệ thống', 'Địa chỉ truy cập', 'Đã kiểm tra đăng nhập']


def rows(accounts, dataset):
    """Rows in staff-list order. Raises if any employee lacks a verified account."""
    positions = {position['code']: position['title'] for position in dataset['positions']}
    missing = [employee['name'] for employee in dataset['employees']
               if not accounts.get(employee['code'], {}).get('verified')]
    if missing:
        raise ValueError('Accounts not verified by a login: %s' % ', '.join(missing))
    result = []
    for number, employee in enumerate(dataset['employees'], start=1):
        account = accounts[employee['code']]
        result.append([number, employee['code'], employee['name'], employee['job_title'],
                       positions.get(employee['position'], ''), account['login'], account['password'],
                       account['role'], SITE, 'Đã đăng nhập thành công'])
    return result


def write(path, table):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Tài khoản'
    sheet.append(['DANH SÁCH TÀI KHOẢN HỆ THỐNG OKR/KPI — PHÒNG KINH DOANH VÀ DỊCH VỤ'])
    sheet.append(['Địa chỉ: %s — Mật khẩu ban đầu cấp riêng từng người; đề nghị đổi mật khẩu sau lần đăng nhập '
                  'đầu tiên (ảnh đại diện › Hồ sơ của tôi › Bảo mật).' % SITE])
    sheet.append([])
    sheet.append(HEADERS)
    for row in table:
        sheet.append(row)
    sheet['A1'].font = Font(bold=True, size=13)
    header_row = 4
    for column in range(1, len(HEADERS) + 1):
        cell = sheet.cell(header_row, column)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F3A5F')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = [5, 9, 24, 34, 40, 28, 16, 22, 24, 22]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = 'A5'
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--accounts', default=str(DEFAULT_ACCOUNTS))
    parser.add_argument('--dataset', default=str(DEFAULT_DATASET))
    parser.add_argument('--out', default=str(DEFAULT_OUT))
    options = parser.parse_args()
    accounts = json.loads(pathlib.Path(options.accounts).read_text(encoding='utf-8'))
    dataset = json.loads(pathlib.Path(options.dataset).read_text(encoding='utf-8'))
    table = rows(accounts, dataset)
    write(pathlib.Path(options.out), table)
    print('%s: %d accounts' % (options.out, len(table)))


if __name__ == '__main__':
    main()
