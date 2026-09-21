# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""The screen walkthrough appended to a deck somebody edits by hand.

The customer read a note that said only what a number was not ("không phải tỷ
lệ hoàn thành dự án") and told us it meant nothing to them. So the rules held
down here are about the writing as much as the file: a note names a figure or
a rule the reader can check, it never opens by denying something, and it is
written in plain words. The mechanics matter too - the appended slides must
take the deck's own typography, keep the page numbering running, and never
reference a screenshot nobody captured.

Run:  python -m unittest discover -s tools/tests -t .
"""
import importlib.util
import pathlib
import re
import tempfile
import unittest

_REPO = pathlib.Path(__file__).resolve().parents[2]


def _load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


tool = _load(_REPO / 'tools' / 'append_kddv_screens.py', 'append_kddv_screens')
HAVE_PPTX = importlib.util.find_spec('pptx') is not None
PNG = bytes.fromhex(
    '89504e470d0a1a0a0000000d4948445200000010000000100806000000'
    '1ff3ff610000001849444154789c6360180560c800130c0d0c0c8c0c0c'
    '0c0c0c0c0cc40001000005b1000130a0d1310000000049454e44ae4260'
    '82')


def notes():
    return [(name, head, body) for name, _title, _lead, items, _where in tool.SCREENS
            for head, body in items]


class WordingCase(unittest.TestCase):

    def test_a_note_never_opens_by_saying_what_something_is_not(self):
        for name, head, body in notes():
            self.assertFalse(body.startswith('Không phải'),
                             f'{name}: "{body}" chỉ phủ định, không giải thích')
            self.assertFalse(head.startswith('Không phải'), name)

    def test_most_screens_carry_a_figure_the_reader_can_check(self):
        with_figure = {name for name, head, body in notes()
                       if re.search(r'\d', head + body)}
        self.assertGreaterEqual(len(with_figure), len(tool.SCREENS) * 2 // 3,
                                'phần lớn màn hình phải kèm một con số cụ thể')

    def test_the_walkthrough_avoids_machine_sounding_words(self):
        self.assertGreater(tool.check_wording(), 0)
        broken = list(tool.SCREENS[:1]) + [
            ('21-cycles', 'Giải pháp toàn diện', 'Một giải pháp tổng thể mạnh mẽ.',
             [('A', 'B')], 'X')]
        with self.assertRaises(SystemExit) as stopped:
            tool.check_wording(broken)
        self.assertIn('toàn diện', str(stopped.exception))

    def test_every_note_is_a_headline_and_a_sentence(self):
        for name, head, body in notes():
            self.assertLessEqual(len(head), 34, f'{name}: tiêu đề ghi chú quá dài')
            self.assertTrue(body.endswith(('.', '?')), f'{name}: "{body}"')
            self.assertGreater(len(body), 20, f'{name}: "{body}" quá ngắn để hiểu')

    def test_each_screen_says_where_it_sits_in_the_menu(self):
        for name, title, lead, items, where in tool.SCREENS:
            self.assertTrue(title and lead and where, name)
            self.assertIn('›', where, f'{name}: phải ghi đường dẫn menu')
            self.assertGreaterEqual(len(items), 2, name)


class PictureCase(unittest.TestCase):

    def test_every_screen_is_one_the_capture_script_takes(self):
        script = (_REPO / 'uat' / 'capture_kddv_guide.mjs').read_text(encoding='utf-8')
        for name, *_rest in tool.SCREENS:
            self.assertIn("shot('%s')" % name, script, name)

    def test_a_missing_screenshot_stops_the_run(self):
        with tempfile.TemporaryDirectory() as empty:
            with self.assertRaises(SystemExit) as stopped:
                tool.check_images(pathlib.Path(empty))
        self.assertIn('Thiếu ảnh', str(stopped.exception))
        self.assertIn('capture_kddv_guide', str(stopped.exception))

    def test_the_screens_are_not_repeated(self):
        names = [name for name, *_rest in tool.SCREENS]
        self.assertEqual(len(names), len(set(names)))


@unittest.skipUnless(HAVE_PPTX, 'cần python-pptx (Python 3.11 ở máy này)')
class DeckCase(unittest.TestCase):

    def setUp(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        root = pathlib.Path(self.folder.name)
        self.images = root / 'img'
        self.images.mkdir()
        for name, *_rest in tool.SCREENS:
            (self.images / f'{name}.png').write_bytes(PNG)
        # A deck in the customer's shape: two slides, their own typography.
        deck = Presentation()
        deck.slide_width, deck.slide_height = 12192000, 6858000
        for index, (title, colour, size) in enumerate(
                [('Bìa', 'FFFFFF', 30.0), ('Tình hình', '16334B', 30.0)]):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for text, box, pt, hexa in (
                    (title, (0.58, 0.40, 12.17, 0.96), size, colour),
                    ('Dẫn nhập', (0.60, 1.40, 12.03, 0.67), 17.25, '637382'),
                    ('%02d' % (index + 1), (12.19, 7.06, 0.62, 0.25), 11.25, '637382'),
                    ('Ghi chú cuối trang', (0.60, 6.55, 12.03, 0.45), 12.75, '637382')):
                shape = slide.shapes.add_textbox(Inches(box[0]), Inches(box[1]),
                                                 Inches(box[2]), Inches(box[3]))
                run = shape.text_frame.paragraphs[0].add_run()
                run.text = text
                run.font.size, run.font.name = Pt(pt), 'Be Vietnam Pro'
                run.font.color.rgb = RGBColor.from_string(hexa)
        self.deck = root / 'deck.pptx'
        deck.save(str(self.deck))

    def test_the_appended_slides_take_the_decks_own_typography(self):
        from pptx import Presentation
        written = tool.append(self.deck, self.images)
        self.assertEqual(written['added'], len(tool.SCREENS) + 1)
        shown = Presentation(str(self.deck))
        self.assertEqual(len(shown.slides), written['total'])
        titles = [shape for shape in shown.slides[2].shapes
                  if shape.has_text_frame and shape.text_frame.text == tool.MAP_SLIDE['title']]
        self.assertEqual(len(titles), 1)
        font = titles[0].text_frame.paragraphs[0].runs[0].font
        self.assertEqual(font.name, 'Be Vietnam Pro', 'phải theo phông của tệp')
        self.assertEqual(str(font.color.rgb), '16334B')

    def test_the_page_numbers_keep_running(self):
        from pptx import Presentation
        tool.append(self.deck, self.images)
        shown = Presentation(str(self.deck))
        numbers = []
        for slide in shown.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and re.fullmatch(r'\d{2}', shape.text_frame.text.strip()):
                    numbers.append(int(shape.text_frame.text))
        self.assertEqual(numbers, list(range(1, len(shown.slides) + 1)))

    def test_every_screen_slide_carries_its_picture_and_menu_path(self):
        from pptx import Presentation
        tool.append(self.deck, self.images)
        shown = Presentation(str(self.deck))
        pictures = sum(1 for slide in shown.slides for shape in slide.shapes
                       if shape.shape_type == 13)
        self.assertEqual(pictures, len(tool.SCREENS))
        words = ' '.join(shape.text_frame.text for slide in shown.slides
                         for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Vị trí trên menu: Hiệu suất › Kế hoạch › Cây liên kết mục tiêu', words)

    def test_a_note_can_be_rewritten_without_losing_its_formatting(self):
        from pptx import Presentation
        changed = tool.retouch(self.deck, {'Dẫn nhập': 'Số liệu đọc tại ngày 21/09/2026.'})
        self.assertEqual(len(changed), 2)
        shown = Presentation(str(self.deck))
        [shape] = [shape for shape in shown.slides[0].shapes
                   if shape.has_text_frame and shape.text_frame.text.startswith('Số liệu')]
        font = shape.text_frame.paragraphs[0].runs[0].font
        self.assertEqual(font.size.pt, 17.25)
        self.assertEqual(str(font.color.rgb), '637382')

    def test_the_whole_deck_can_be_put_on_a_font_the_reader_has(self):
        """A font nobody has installed is substituted when the file opens, and
        every box measured here stops fitting there."""
        from pptx import Presentation
        tool.append(self.deck, self.images)
        touched = tool.set_font(self.deck, 'Segoe UI')
        self.assertGreater(touched, 50)
        shown = Presentation(str(self.deck))
        names = {run.font.name for slide in shown.slides for shape in slide.shapes
                 if shape.has_text_frame
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs}
        self.assertEqual(names, {'Segoe UI'})

    def test_rewriting_text_that_is_not_there_stops_the_run(self):
        with self.assertRaises(SystemExit) as stopped:
            tool.retouch(self.deck, {'Không có câu này': 'gì đó'})
        self.assertIn('Không tìm thấy', str(stopped.exception))


if __name__ == '__main__':
    unittest.main()
