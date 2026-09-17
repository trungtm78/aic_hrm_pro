# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""The customer account sheet only lists accounts proven by a login."""
import importlib.util
import pathlib
import tempfile
import unittest

_REPO = pathlib.Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location('export_kddv_accounts', _REPO / 'tools' / 'export_kddv_accounts.py')
export = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(export)

DATASET = {
    'positions': [{'code': 'CV', 'title': 'Chuyên viên'}],
    'employees': [
        {'code': 'E1', 'name': 'An', 'job_title': 'Chuyên viên', 'position': 'CV'},
        {'code': 'E2', 'name': 'Bình', 'job_title': 'Chuyên viên', 'position': 'CV'},
    ],
}


def _account(code, verified=True):
    return {'code': code, 'login': '%s@example.com' % code, 'password': 'Pw-%s' % code,
            'role': 'Hiệu suất: Người dùng', 'verified': verified}


class ExportCase(unittest.TestCase):

    def test_rows_follow_the_staff_order(self):
        table = export.rows({'E2': _account('E2'), 'E1': _account('E1')}, DATASET)
        self.assertEqual([row[1] for row in table], ['E1', 'E2'])
        self.assertEqual(table[0][5:8], ['E1@example.com', 'Pw-E1', 'Hiệu suất: Người dùng'])

    def test_unverified_or_missing_account_blocks_the_sheet(self):
        with self.assertRaisesRegex(ValueError, 'Bình'):
            export.rows({'E1': _account('E1')}, DATASET)
        with self.assertRaisesRegex(ValueError, 'An'):
            export.rows({'E1': _account('E1', verified=False), 'E2': _account('E2')}, DATASET)

    def test_workbook_is_written(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as folder:
            path = pathlib.Path(folder) / 'accounts.xlsx'
            export.write(path, export.rows({'E1': _account('E1'), 'E2': _account('E2')}, DATASET))
            sheet = openpyxl.load_workbook(path).active
            self.assertEqual(sheet.cell(4, 1).value, 'STT')
            self.assertEqual(sheet.cell(6, 3).value, 'Bình')


if __name__ == '__main__':
    unittest.main()
