# -*- coding: utf-8 -*-
# Part of AIC HRM Pro. See LICENSE file for full copyright and licensing details.
"""The progress deck: what it says, and that it cannot say it wrongly.

This deck is presented to the customer's leadership, so the failure that
matters is not a crash - it is a slide that looks right and is wrong: a
number typed instead of read, a month with no figures drawn as a zero, a
screenshot silently missing, a screen named that the customer does not have.
Those are what the cases below hold down. The slide content is built without
python-pptx, matplotlib or reportlab, so these checks run on any interpreter;
only the three file-writing cases need the 3.11 interpreter.

Run:  python -m unittest discover -s tools/tests -t .
"""
import copy
import importlib.util
import pathlib
import re
import tempfile
import unittest

_REPO = pathlib.Path(__file__).resolve().parents[2]


def _load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


deck_tool = _load(_REPO / 'tools' / 'build_kddv_deck.py', 'build_kddv_deck')
_handbook_tests = _load(pathlib.Path(__file__).with_name('test_build_kddv_handbook.py'),
                        'test_build_kddv_handbook')


def _have(name):
    return importlib.util.find_spec(name) is not None


HAVE_PPTX, HAVE_PDF, HAVE_CHARTS = _have('pptx'), _have('reportlab'), _have('matplotlib')
NEEDS_311 = ('cần trình thông dịch có python-pptx, matplotlib và reportlab (Python 3.11 ở máy này)')
PNG = _handbook_tests.PNG


def dataset():
    """The handbook's description of production, plus the two fields the deck
    reads that the handbook does not."""
    data = copy.deepcopy(_handbook_tests.dataset())
    for cards in data['cards'].values():
        for card in cards:
            card['rag'] = ('green' if card['data_coverage'] and card['score_covered'] >= 0.7
                           else 'amber' if card['data_coverage'] else 'none')
    # The roll-up view reads the same cards, so its figures have to agree
    # with them: two of the three people have figures, and their measured
    # score is (100% + 60,4%) / 2. A fixture where the two disagree would
    # hide the very mistake this deck was built to stop.
    measured = [card for card in data['cards'][7] if card['data_coverage'] > 0]
    for row in data['department_report']:
        row['measured_employee_count'] = len(measured)
        row['avg_score_covered'] = (sum(card['score_covered'] for card in measured)
                                    / len(measured))
    return data


def source():
    """The file reader's dataset, plus the monthly totals the deck reconciles
    against (the handbook does its reconciliation per stream instead)."""
    data = copy.deepcopy(_handbook_tests.source())
    data['revenue'] = {'7': {'total': 38.03, 'has_data': True},
                       '8': {'total': 43.99, 'has_data': True},
                       '9': {'total': 0.0, 'has_data': False}}
    # The file reader writes a block per month, not a bare number.
    data['cost'] = {'7': {'total': 18.445937012, 'groups': []}}
    return data


ENGAGEMENT = {
    'meetings': [{'date': '05/08/2026', 'title': 'Buổi 1 — Khảo sát',
                  'topics': ['Cách giao KPI hiện tại', 'Nguồn số liệu doanh thu']}],
    'requirements': [{'need': 'Chấm KPI theo tháng', 'answer': 'Bảng điểm cá nhân theo chu kỳ tháng'}],
}

DATA = dataset()
SOURCE = source()
FACTS = deck_tool.facts(DATA, SOURCE, ENGAGEMENT)
DECK = deck_tool.deck(FACTS)


def build(data=None, engagement=ENGAGEMENT):
    return deck_tool.deck(deck_tool.facts(data or dataset(), source(), engagement))


def texts(slides):
    return deck_tool.visible_strings(slides)


class PaletteCase(unittest.TestCase):
    """The deck must look like the system it reports on."""

    def test_the_design_tokens_convert_to_the_colours_on_screen(self):
        self.assertEqual(deck_tool.oklch_to_rgb(0.45, 0.14, 235.0), '005E97')
        self.assertEqual(deck_tool.oklch_to_rgb(0.60, 0.13, 150.0), '3B9555')
        self.assertEqual(deck_tool.oklch_to_rgb(0.55, 0.18, 25.0), 'C53637')

    def test_a_colour_outside_the_screen_gamut_still_gives_a_usable_value(self):
        self.assertRegex(deck_tool.oklch_to_rgb(0.99, 0.40, 150.0), r'^[0-9A-F]{6}$')


class PartnerNameCase(unittest.TestCase):
    """A chart label that runs off the picture loses the company's name."""

    def test_a_long_legal_name_is_shortened_but_still_recognisable(self):
        short = deck_tool.short_name('Công ty Cổ phần viễn thông FPT')
        self.assertEqual(short, 'CTCP viễn thông FPT')
        long_one = deck_tool.short_name('CÔNG TY CỔ PHẦN TECHNOLOGY CONVERGENCE CORPORATION')
        self.assertLessEqual(len(long_one), 34)
        self.assertTrue(long_one.startswith('CTCP TECHNOLOGY'), long_one)

    def test_a_short_name_is_left_alone(self):
        self.assertEqual(deck_tool.short_name('Apple'), 'Apple')

    def test_the_chart_labels_fit_the_picture(self):
        chart = next(c for c in deck_tool.charts_of(DECK) if c.key == 'doanh-thu-theo-doi-tac')
        for label in chart.categories:
            self.assertLessEqual(len(label), 34, label)


class OutlineCase(unittest.TestCase):

    def test_the_deck_opens_on_a_cover_and_ends_with_the_appendix(self):
        self.assertEqual(DECK[0].kind, 'cover')
        self.assertEqual(DECK[1].kind, 'kpis')
        self.assertEqual(DECK[-1].kind, 'table')
        self.assertGreaterEqual(len(DECK), 30)

    def test_the_sections_follow_the_order_the_customer_asked_for(self):
        self.assertEqual(
            [slide.title for slide in DECK if slide.kind == 'section'],
            ['1. Tiếp nhận yêu cầu',
             '2. Chức năng đã đưa vào vận hành',
             '3. Dữ liệu thật đã đưa vào hệ thống',
             '4. Với dữ liệu này, hệ thống đọc ra được gì',
             '5. Hệ thống chấm điểm từng người như thế nào',
             '6. Vì sao tin được những con số này',
             '7. Năng lực hệ thống và hiệu quả mang lại',
             'Phụ lục'])

    def test_every_slide_says_what_it_is_and_why_it_is_there(self):
        for slide in DECK:
            self.assertIn(slide.kind, deck_tool.KINDS, slide.title)
            self.assertTrue(slide.title, slide.kind)
            if slide.kind not in ('section', 'cover'):
                self.assertTrue(slide.lead, slide.title)

    def test_every_kind_can_be_drawn_in_both_outputs(self):
        used = {slide.kind for slide in DECK}
        for handlers in (deck_tool.PPTX_HANDLERS, deck_tool.PDF_HANDLERS):
            self.assertEqual(used - set(handlers), set())

    def test_no_table_is_taller_than_a_slide(self):
        for slide in DECK:
            if slide.table:
                self.assertLessEqual(len(slide.table.rows), deck_tool.MAX_TABLE_ROWS, slide.title)
                for row in slide.table.rows:
                    self.assertEqual(len(row), len(slide.table.headers), slide.title)


class NumbersCase(unittest.TestCase):
    """Every figure is read from the system, not typed into the deck."""

    def test_the_revenue_headline_follows_the_ledger(self):
        changed = dataset()
        changed['ledger_by_account'][(7, '51131')] = 99_000_000_000.0
        expected = deck_tool.vn(
            sum(changed['ledger_by_account'].values()) / 1e9, 1) + ' tỷ'
        moved = [kpi.value for slide in build(changed) for kpi in slide.kpis]
        self.assertIn(expected, moved)
        self.assertNotIn(expected, [kpi.value for slide in DECK for kpi in slide.kpis])

    def test_the_objective_score_is_weighted_by_objective_weight(self):
        changed = dataset()
        changed['objectives'].append({'id': 2, 'code': 'O2', 'name': 'MAU', 'weight': 50.0,
                                      'score': 0.0, 'score_covered': 0.0, 'data_coverage': 0.0,
                                      'employee_id': [5, 'TP']})
        values = [kpi.value for slide in build(changed) for kpi in slide.kpis]
        self.assertIn('35%', values, 'điểm 70% với trọng số 50/100 phải ra 35%')

    def test_a_month_with_no_figures_is_stated_as_such_not_as_zero(self):
        rows = [row for slide in DECK if slide.table
                for row in slide.table.rows if row and row[0] == 'Tháng 8']
        self.assertTrue(rows)
        self.assertIn('Chưa có số liệu', rows[0])

    def test_the_share_the_system_scores_by_itself_is_computed_not_written(self):
        changed = dataset()
        for target in changed['targets'].values():
            target['metric_source_id'] = False
        without = ' '.join(texts(build(changed)))
        self.assertIn('0,0%', without)
        self.assertNotEqual(deck_tool.facts(changed, source())['automation']['share'],
                            FACTS['automation']['share'])

    def test_the_cost_headline_counts_only_the_cost_entries(self):
        """The accrual sits in the same list of entries; counted with the
        cost it claims the cost came from one more document than it did."""
        changed = dataset()
        changed['entries'].append({'ref': 'DTHU2026-T08', 'name': 'DTHU/2026/08/0001',
                                   'date': '2026-08-31', 'journal_id': [2, 'Dự thu doanh thu'],
                                   'amount_total': 22000000000.0})
        bundle = deck_tool.facts(changed, source())
        self.assertEqual(bundle['cost_entries'], 1)
        self.assertEqual(bundle['accrual_entries'], 1)
        notes = [kpi.note for slide in build(changed) for kpi in slide.kpis]
        self.assertIn('1 bút toán chi phí', notes)

    def test_the_example_scorecards_are_read_from_the_cards(self):
        examples = FACTS['examples']
        self.assertTrue(examples)
        for example in examples:
            self.assertGreater(example['coverage'], 0)
            self.assertEqual(sum(line['weight'] for line in example['lines']),
                             sum(line['weight'] for line in DATA['lines'][
                                 next(card['id'] for card in DATA['cards'][7]
                                      if card['employee_id'][1] == example['name'])]))

    def test_an_example_line_carries_the_unit_of_its_target(self):
        rows = [row for slide in DECK if slide.table and 'Chỉ tiêu giao' in slide.table.headers
                for row in slide.table.rows]
        self.assertTrue(rows)
        column = next(slide.table.headers.index('Chỉ tiêu giao') for slide in DECK
                      if slide.table and 'Chỉ tiêu giao' in slide.table.headers)
        self.assertTrue(any('tỷ' in row[column] or '%' in row[column] for row in rows),
                        'người đọc phải biết 49,15 là tỷ đồng hay hợp đồng')

    def test_an_example_line_says_where_its_figure_comes_from(self):
        sources = {line['source'] for example in FACTS['examples'] for line in example['lines']}
        self.assertLessEqual(sources, {deck_tool.AUTOMATIC, deck_tool.BY_HAND})
        self.assertIn(deck_tool.AUTOMATIC, sources)

    def test_a_line_nobody_reported_says_so_instead_of_showing_a_zero(self):
        """A reported 0 and a figure nobody has entered look identical on a
        printed slide, and the second one is what gets read as failure."""
        checked = 0
        for slide in DECK:
            if slide.table and 'Thực hiện' in slide.table.headers:
                actual = slide.table.headers.index('Thực hiện')
                achieved = slide.table.headers.index('Đạt')
                for row in slide.table.rows:
                    if row[achieved] == '—':
                        self.assertEqual(row[actual], 'Chưa có số liệu', slide.title)
                        checked += 1
        self.assertGreater(checked, 0, 'ví dụ phải có dòng chưa có số liệu')


class TextCase(unittest.TestCase):

    FORBIDDEN = ('aic.hrm', 'account.move', 'search_read', 'None', 'nan', 'inf',
                 '{', '}', '%s', 'Traceback', 'cycle_id')

    def test_nothing_technical_or_unformatted_reaches_the_customer(self):
        for text in texts(DECK):
            for forbidden in self.FORBIDDEN:
                self.assertNotIn(forbidden, text, text)

    def test_the_deck_is_written_in_vietnamese(self):
        vietnamese = re.compile(r'[àáâãèéêìíòóôõùúýăđĩũơưạ-ỹÀ-Ỹ]')
        for slide in DECK:
            self.assertRegex(slide.title, vietnamese, slide.title)

    def test_numbers_are_written_the_vietnamese_way(self):
        for slide in DECK:
            for row in (slide.table.rows if slide.table else ()):
                for value in row:
                    self.assertIsInstance(value, str)
                    if re.fullmatch(r'-?\d+\.\d+', value):
                        self.fail('số theo kiểu Anh trong "%s": %r' % (slide.title, value))

    def test_the_measured_score_is_written_the_same_way_on_every_slide(self):
        """82,6% on one slide and 83% on another reads as two measurements of
        the same thing, and a reader who spots it stops trusting both."""
        column = 'Điểm trên phần có số liệu'
        shown = set()
        for slide in DECK:
            table = slide.table
            if table and column in table.headers:
                index = table.headers.index(column)
                for row in table.rows:
                    if row[0].startswith('Tháng 7'):
                        shown.add(row[index])
        self.assertEqual(shown, {deck_tool.percent(FACTS['months'][0]['score_covered'])},
                         'cùng một con số phải viết giống nhau ở mọi trang')

    def test_the_meetings_slide_says_when_it_has_nothing_to_show(self):
        empty = build(engagement={})
        rows = [row for slide in empty if slide.table and 'Buổi làm việc' in slide.table.headers
                for row in slide.table.rows]
        self.assertTrue(any('Chưa có thông tin' in value for row in rows for value in row))


class ChartCase(unittest.TestCase):

    def test_every_chart_has_numbers_behind_every_category(self):
        for chart in deck_tool.charts_of(DECK):
            self.assertTrue(chart.categories, chart.key)
            for series in chart.series:
                self.assertEqual(len(series.values), len(chart.categories), chart.key)
                self.assertTrue(any(value is not None for value in series.values), chart.key)
                for value in series.values:
                    self.assertTrue(value is None or isinstance(value, float), chart.key)

    def test_a_month_without_figures_is_a_gap_in_the_chart_not_a_zero(self):
        chart = next(c for c in deck_tool.charts_of(DECK) if c.key == 'ke-hoach-va-thuc-hien-quy-3')
        actual = next(series for series in chart.series if series.name == 'Thực hiện')
        self.assertIsNone(actual.values[-1], 'tháng 9 chưa có số liệu')

    def test_an_unreported_objective_is_a_gap_not_a_nought(self):
        changed = dataset()
        changed['objectives'].append({'id': 3, 'code': 'O3', 'name': 'Quy trình', 'weight': 20.0,
                                      'score': 0.0, 'score_covered': 0.0, 'data_coverage': 0.0,
                                      'employee_id': [5, 'TP']})
        chart = next(c for c in deck_tool.charts_of(build(changed)) if c.key == 'muc-tieu-quy-3')
        self.assertIsNone(chart.series[0].values[-1])

    def test_chart_files_have_safe_distinct_names(self):
        keys = [chart.key for chart in deck_tool.charts_of(DECK)]
        self.assertEqual(len(keys), len(set(keys)))
        for key in keys:
            self.assertRegex(key, r'^[a-z0-9-]+$', key)

    def test_a_chart_kind_the_slides_use_is_a_kind_a_renderer_knows(self):
        for chart in deck_tool.charts_of(DECK):
            self.assertIn(chart.kind, ('column', 'stacked', 'bar', 'line', 'combo', 'doughnut'))
            if chart.kind == 'combo':
                self.assertFalse(chart.native, 'PowerPoint không vẽ được cột kèm đường')


class PictureCase(unittest.TestCase):

    def test_a_missing_screenshot_stops_the_build(self):
        with tempfile.TemporaryDirectory() as empty:
            with self.assertRaises(SystemExit) as stopped:
                deck_tool.check_pictures(DECK, pathlib.Path(empty))
        self.assertIn('Thiếu ảnh', str(stopped.exception))
        self.assertIn('capture_kddv_guide', str(stopped.exception))

    def test_every_picture_is_one_the_capture_script_takes(self):
        script = (_REPO / 'uat' / 'capture_kddv_guide.mjs').read_text(encoding='utf-8')
        for picture in deck_tool.pictures_of(DECK):
            self.assertIn("shot('%s')" % picture.name, script, picture.name)
            self.assertTrue(picture.caption, picture.name)

    def test_the_deck_claims_no_screen_without_a_caption_for_it(self):
        for _group, entries in deck_tool.FUNCTIONS:
            for name, note in entries:
                self.assertTrue(name and note)


class DependencyCase(unittest.TestCase):

    def test_the_slides_are_built_without_any_rendering_library(self):
        """So the numbers can be checked on any interpreter, and a missing
        package can never change what the deck says."""
        source_text = (_REPO / 'tools' / 'build_kddv_deck.py').read_text(encoding='utf-8')
        header = source_text.split('def _load(', 1)[0]
        for library in ('pptx', 'matplotlib', 'reportlab', 'PIL'):
            self.assertNotIn('import %s' % library, header, library)


@unittest.skipUnless(HAVE_CHARTS and HAVE_PPTX and HAVE_PDF, NEEDS_311)
class ArtifactCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        root = pathlib.Path(cls.folder.name)
        cls.images = root / 'img'
        cls.images.mkdir()
        for picture in deck_tool.pictures_of(DECK):
            (cls.images / f'{picture.name}.png').write_bytes(PNG)
        cls.charts = deck_tool.write_charts(DECK, root / 'charts')
        cls.pptx = root / 'deck.pptx'
        deck_tool.render_pptx(DECK, cls.pptx, cls.charts, cls.images)
        cls.pdf = root / 'deck.pdf'
        cls.pages = deck_tool.render_pdf(DECK, cls.pdf, cls.charts, cls.images)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def test_one_png_per_chart_and_none_of_them_empty(self):
        self.assertEqual(set(self.charts), {chart.key for chart in deck_tool.charts_of(DECK)})
        for path in self.charts.values():
            self.assertGreater(path.stat().st_size, 5000, path.name)

    def test_the_powerpoint_is_widescreen_and_holds_every_slide(self):
        from pptx import Presentation
        shown = Presentation(str(self.pptx))
        self.assertEqual(len(shown.slides), len(DECK))
        self.assertEqual((shown.slide_width, shown.slide_height), deck_tool.SLIDE_EMU)

    def test_the_charts_the_customer_can_edit_are_really_charts(self):
        from pptx import Presentation
        shown = Presentation(str(self.pptx))
        native = [shape.chart for slide in shown.slides for shape in slide.shapes
                  if shape.has_chart]
        self.assertEqual(len(native), sum(1 for chart in deck_tool.charts_of(DECK) if chart.native))
        wanted = next(chart for chart in deck_tool.charts_of(DECK) if chart.native)
        self.assertEqual(list(native[0].plots[0].categories), list(wanted.categories))

    def test_the_headline_figures_really_are_on_a_slide(self):
        from pptx import Presentation
        shown = Presentation(str(self.pptx))
        words = ' '.join(shape.text_frame.text for slide in shown.slides
                         for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Kinh doanh và Dịch vụ', words)
        self.assertIn(deck_tool.percent(FACTS['okr_score'], 0), words)

    def test_the_pdf_has_exactly_one_page_per_slide(self):
        self.assertEqual(self.pages, len(DECK))
        self.assertTrue(self.pdf.read_bytes().startswith(b'%PDF'))
        self.assertGreater(self.pdf.stat().st_size, 20000)


if __name__ == '__main__':
    unittest.main()
